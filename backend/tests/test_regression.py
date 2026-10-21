"""Offline regression checks; run with python -m tests.run_checks."""
import os
import atexit
import tempfile
from pathlib import Path
_workspace = tempfile.TemporaryDirectory(prefix="claims-regression-")
atexit.register(_workspace.cleanup)
os.environ["DATABASE_URL"] = "sqlite:///" + (Path(_workspace.name) / "claims.sqlite").as_posix()
os.environ["UPLOAD_DIR"] = str(Path(_workspace.name) / "uploads")
os.environ["DISABLE_AUTH"] = "false"
os.environ["APP_SECRET"] = "isolated-regression-secret-not-for-production"
for key in ("WDT_SID", "WDT_APPKEY", "WDT_APPSECRET", "WDT_QIMEN_APPKEY", "WDT_QIMEN_APPSECRET", "WDT_QIMEN_TARGET_APPKEY"):
    os.environ[key] = ""
import io
import json
from datetime import date, datetime
from decimal import Decimal
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select
from app import api
from app.database import Base, engine, SessionLocal
from app.models import ClaimEvent, ClaimStatus, AfterSalesType, Evidence, DutyParty
from app.security import require_auth
from app.services.engine import detect_original_return, serialize_event, summarize
from app.services.wangdian import WangdianClient
from app.services.report import build_report_workbook
from app.services.legacy_import import import_legacy
from openpyxl import load_workbook

atexit.register(engine.dispose)
Base.metadata.create_all(engine)
app=FastAPI();app.include_router(api.router)
app.dependency_overrides[require_auth]=lambda:'isolated-audit'
client=TestClient(app,raise_server_exceptions=False)
results=[]
def record(name,expected,actual):
    results.append({'name':name,'expected':expected,'actual':actual,'pass':expected==actual})
def verified(*_):
    return {'ok':True,'shipment':{'source_tid':'WDT-SOURCE','logistics_no':'TRACK-SOURCE','shop_name':'Test','goods_total_amount':'59.90'}}
api._lookup_shipment=verified
base={'event_date':'2026-10-15T10:15','incident_type':'damage','claim_amount':'20','wdt_tid':'MANUAL'}
for key in ('incident_type','event_date','claim_amount'):
    p=dict(base);p.pop(key);record('new_required_'+key,400,client.post('/api/events',json=p).status_code)
api._lookup_shipment=lambda *_:{'ok':True,'shipment':{'source_tid':''}}
record('new_requires_original_order',409,client.post('/api/events',json=base).status_code)
api._lookup_shipment=verified
r=client.post('/api/events',json=base)
record('new_uses_verified_source','WDT-SOURCE',r.json().get('event',{}).get('wdt_tid'))
record('new_blocks_existing_duplicate',409,client.post('/api/events',json=base).status_code)

db=SessionLocal()
def seed(code,**kwargs):
    defaults=dict(event_date=date(2026,10,15),event_time=datetime(2026,10,15,10,15),wdt_tid=code,ship_logistics_no='TRACK-'+code,claim_amount=Decimal('20'))
    defaults.update(kwargs)
    e=ClaimEvent(code=code,**defaults);db.add(e);db.commit();return e
seed('EDIT-ID')
record('edit_cannot_clear_original',True,client.put('/api/events/EDIT-ID',json={'wdt_tid':''}).status_code in (400,409))
seed('EDIT-DUP')
record('edit_blocks_duplicate',409,client.put('/api/events/EDIT-DUP',json={'wdt_tid':'WDT-SOURCE'}).status_code)
seed('NEGATIVE')
record('reject_negative_claim',400,client.put('/api/events/NEGATIVE',json={'claim_amount':'-10'}).status_code)
seed('LOWER',claim_amount=Decimal('100'),received_amount=Decimal('80'))
record('lower_claim_must_check_received',400,client.put('/api/events/LOWER',json={'claim_amount':'20'}).status_code)
seed('ZERO')
record('zero_claim_cannot_receive_more',400,client.put('/api/events/ZERO',json={'claim_amount':'0','received_amount':'50'}).status_code)
seed('PAID')
record('paid_requires_actual_receipt',400,client.put('/api/events/PAID',json={'status':'paid','received_amount':'0'}).status_code)
seed('DUTY-NO-BASIS')
record('duty_without_basis_returns_validation',400,client.put('/api/events/DUTY-NO-BASIS',json={'duty_party':'express'}).status_code)
seed('DUTY-WITH-BASIS')
r=client.put('/api/events/DUTY-WITH-BASIS',json={'duty_party':'warehouse','duty_basis':'Isolated test evidence checked'})
record('duty_with_basis_saves',200,r.status_code)
record('duty_change_history',1,len(client.get('/api/events/DUTY-WITH-BASIS').json().get('duty_changes',[])))
seed('TIME')
client.put('/api/events/TIME',json={'event_date':'2026-10-15T11:45'})
logs=client.get('/api/events/TIME/audit').json()['rows']
record('minute_change_audited',True,any('11:45' in x.get('new_value','') for x in logs))
seed('SORT-EARLY-ID',event_time=datetime(2026,10,15,23,0),shop_name='SORT-TEST')
seed('SORT-LATE-ID',event_time=datetime(2026,10,15,8,0),shop_name='SORT-TEST')
record('sort_by_registration_minute','SORT-EARLY-ID',client.get('/api/events?shop=SORT-TEST').json()['rows'][0]['code'])
record('end_date_excludes_next_day',0,client.get('/api/events?start=2026-10-14&end=2026-10-14').json()['total'])
seed('SEPT',event_date=date(2026,9,30),event_time=datetime(2026,9,30,23,59),shop_name='MONTH-TEST')
seed('OCT',event_date=date(2026,10,1),event_time=datetime(2026,10,1,0,1),shop_name='MONTH-TEST')
record('monthly_boundary',['OCT'],[x['code'] for x in client.get('/api/events?month=2026-10&shop=MONTH-TEST').json()['rows']])
returned=seed('NORMAL-RETURN',after_sales_type=AfterSalesType.RETURN_REFUND,stockin_found=True)
record('warehouse_receipt_does_not_prove_original_return',False,detect_original_return(returned)['suggested_type']=='original_return')

# Test matching using deterministic upstream replies; no external requests.
wdt=object.__new__(WangdianClient)
wdt._request=lambda *_:{'stockin_list':[{'src_order_no':'UNRELATED','remark':'ORDER-91234','warehouse_name':'Warehouse','status':'80'}]}
record('stockin_blank_keys_no_global_match',0,len(wdt.search_refund_stockins(order_no='',logistics_no='')))
record('stockin_requires_exact_order_match',0,len(wdt.search_refund_stockins(order_no='123',logistics_no='')))
wdt._request=lambda *_:{'trades':[{'logistics_no':'TRACK-A','src_tids':'ORDER-A'}]}
record('both_keys_must_agree',0,len(wdt.find_sales_stockouts(logistics_no='TRACK-A',src_tid='ORDER-B')))

report=load_workbook(io.BytesIO(build_report_workbook([db.get(ClaimEvent,returned.id)],summarize([returned]))))
record('xlsx_contains_minute',True,'10:15' in str(report['明细'].cell(2,2).value))
record('xlsx_contains_after_sales_type',True,'退回退款' in str(list(report['明细'].values)))
record('order_amount_serialized','59.90',client.get('/api/events/WDT-SOURCE').json()['order_amount'])
record('shipment_goods_kept_in_detail',True,'goods_items' in client.get('/api/events/WDT-SOURCE').json())
record('blank_claim_update_rejected',400,client.put('/api/events/WDT-SOURCE',json={'claim_amount':''}).status_code)
record('nonfinite_claim_rejected',400,client.put('/api/events/WDT-SOURCE',json={'claim_amount':'NaN'}).status_code)
record('subcent_amount_rejected',400,client.put('/api/events/WDT-SOURCE',json={'claim_amount':'1.001'}).status_code)
record('optional_evidence_not_required',0,len(client.get('/api/events/WDT-SOURCE').json()['evidence']))
record('communication_merchant',200,client.post('/api/events/WDT-SOURCE/replies',json={'source':'owner','content':'Merchant statement'}).status_code)
record('communication_warehouse',200,client.post('/api/events/WDT-SOURCE/replies',json={'source':'warehouse','content':'Warehouse response'}).status_code)
record('communication_both_sides',['owner','warehouse'],[r['source'] for r in client.get('/api/events/WDT-SOURCE').json()['replies']])
snapshot_row={'src_trade_no':'ORDER-X','details_list':[{'spec_no':'SKU1','goods_name':'Test goods','goods_count':2,'goods_amount':'40'}],'goods_list':[{'spec_no':'SKU1','goods_name':'Test goods','goods_count':2,'goods_amount':'40'}]}
record('goods_alias_not_double_counted','40',str(WangdianClient.parse_sales_stockout(snapshot_row)['goods_total_amount']))
wb=load_workbook(io.BytesIO(build_report_workbook([db.get(ClaimEvent,1)],summarize([db.get(ClaimEvent,1)]))))
header=[c.value for c in wb['明细'][1]]
wb['明细'].cell(2,header.index('实收金额')+1).value=5
output=io.BytesIO();wb.save(output)
rr=client.post('/api/report/import',files={'file':('report.xlsx',output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
record('xlsx_roundtrip_import',200,rr.status_code)
record('xlsx_roundtrip_actual_receipt','5.00',client.get('/api/events/WDT-SOURCE').json()['received_amount'])
legacy='日期,订单号,快递单号,售后类型,赔付金额\n2026-10-15,UNVERIFIED,T-1,破损,10\n'.encode('utf-8-sig')
result=import_legacy(db,legacy,verify_shipment=lambda *_:{'ok':False})
record('legacy_import_requires_wdt',0,result['created'])
print(json.dumps({'total':len(results),'passed':sum(x['pass'] for x in results),'checks':results},ensure_ascii=False,indent=2))
assert all(x['pass'] for x in results), 'Some logic checks failed'
