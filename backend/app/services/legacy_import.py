"""旧 Excel 台账导入。

导入原则（不捏造）：
- 台账「赔付金额」= 申请金额（用户已确认口径）→ 落 claim_amount
- 「实际到手」若有值 → received_amount；没有 → 留 0 等确认
- 顾客已赔（refund_amount）未知 → 留 0，后续用旺店通同步回填
- 回复列按表头出现顺序识别：第1组→仓库回复，第2组→货主备注；=DISPIMG 公式跳过
"""
from __future__ import annotations

import csv
import io
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import ClaimEvent, ClaimReply, ClaimStatus, DutyParty, ReplySource, platform_of

HEADERS = {
    "日期": "date", "店铺名称": "shop", "订单号": "order_no",
    "快递公司": "carrier", "快递单号": "logistics_no",
    "售后类型": "type", "售后内容": "claim",
    "赔付金额": "claim_amount", "实际到手": "received",
    "责任方": "duty",
    "备注": "note", "事件编号": "code", "状态": "status",
}
# 表头别名：新模板把该列写作「仓库赔付金额」，两种写法都接受，旧台账文件继续可用。
HEADER_ALIASES = {"仓库赔付金额": "赔付金额"}
DUTY_MAP = {"仓库": DutyParty.WAREHOUSE, "仓储": DutyParty.WAREHOUSE,
            "仓库方": DutyParty.WAREHOUSE, "快递": DutyParty.EXPRESS, "承运": DutyParty.EXPRESS}
TYPE_MAP = {"破损": "damage", "损坏": "damage", "丢件": "loss", "丢失": "loss",
            "短少": "shortage", "少件": "shortage", "错发": "shortage",
            "漏发": "shortage", "漏发/错发": "shortage", "错发/漏发": "shortage",
            "未派送/派送错误": "undelivered"}

STATUS_MAP = {
    "待提交": ClaimStatus.DRAFT, "待更新数据": ClaimStatus.DRAFT, "已提交": ClaimStatus.SUBMITTED,
    "已赔付": ClaimStatus.PAID, "已驳回": ClaimStatus.REJECTED,
    "仲裁驳回": ClaimStatus.REJECTED, "无法追责": ClaimStatus.REJECTED,
    "已放弃": ClaimStatus.GAVE_UP,
}


def _money(v: Any) -> Decimal | None:
    s = str(v or "").strip()
    if not s:
        return None
    if "-" in s:
        raise ValueError("导入金额不能为负数")
    digits = "".join(ch for ch in s if ch.isdigit() or ch == ".")
    if not digits:
        return None
    try:
        return Decimal(digits)
    except InvalidOperation:
        return None


def _alias(row: dict[str, Any]) -> dict[str, Any]:
    """把新模板的表头别名（如「仓库赔付金额」）归一成下游使用的旧名。"""
    for alias, canonical in HEADER_ALIASES.items():
        if alias in row and canonical not in row:
            row[canonical] = row[alias]
    return row


def _rows(content: bytes) -> list[dict[str, Any]]:
    if not content.startswith(b"PK"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("gb18030")
        return [_alias(r) for r in csv.DictReader(io.StringIO(text))]
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        raw = list(ws.iter_rows(values_only=True))
        wb.close()
        if not raw:
            return []
        header = [str(c).strip() if c is not None else "" for c in raw[0]]
        out = []
        for r in raw[1:]:
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            out.append({header[i]: r[i] for i in range(min(len(header), len(r)))})
        return [_alias(r) for r in out]
    except ImportError:
        text = content.decode("utf-8-sig", "replace")
        return [_alias(r) for r in csv.DictReader(io.StringIO(text))]


def _parse_date(v: Any) -> Any:
    """返回 date 或抛 ValueError。

    支持：datetime/date 对象、2026-08-01、2026-08-01 00:00:00、8月12日、2026年8月12日。
    Excel 日期单元格经 openpyxl 读出的是 datetime 对象，必须直接取 .date()。
    """
    import datetime as _dt

    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    s = str(v or "").strip()
    if not s:
        raise ValueError("空日期")
    s = s.replace("年", "-").replace("月", "-").replace("日", "")
    # 去掉时间部分：2026-09-01 00:00:00 / 2026-09-01T00:00:00
    s = s.split(" ")[0].split("T")[0].strip()
    parts = [p for p in s.split("-") if p.strip()]
    today = _dt.date.today()
    try:
        if len(parts) == 2:
            m, d = int(parts[0]), int(parts[1])
            cand = _dt.date(today.year, m, d)
            return cand if cand <= today else _dt.date(today.year - 1, m, d)
        if len(parts) == 3:
            return _dt.date(int(parts[0]), int(parts[1]), int(parts[2]))
    except (ValueError, IndexError):
        pass
    raise ValueError(f"无法识别: {v}")


def import_legacy(db: Session, content: bytes, *, verify_shipment=None, apply_shipment=None, validate_amounts=None) -> dict[str, Any]:
    rows = _rows(content)
    created = updated = skipped = 0
    issues: list[str] = []
    skipped_rows: list[dict] = []

    def _skip(row_meta: dict, reason: str) -> None:
        skipped_rows.append({**row_meta, "reason": reason})
        issues.append(f"第{row_meta['row']}行：{reason}")
        nonlocal skipped
        skipped += 1

    for i, raw in enumerate(rows, start=2):
        get = lambda k: raw.get(k, "")  # noqa: E731

        order_no = str(get("订单号") or "").strip()
        code = str(get("事件编号") or "").strip()
        tracking = str(get("快递单号") or "").strip()
        row_meta = {
            "row": i, "code": code, "order_no": order_no,
            "logistics_no": tracking,
            "shop": str(get("店铺名称") or ""),
            "date": str(get("日期") or ""),
            "type": str(get("售后类型") or ""),
            "amount": str(get("赔付金额") or ""),
        }
        # 旧台账（如「聚小翻9月售后汇总」）常把单号统一记在「快递单号」列，
        # 订单号与发货物流单号混排，故三者为任一有值即可继续。
        if not code and not order_no and not tracking:
            _skip(row_meta, "既无「事件编号」「订单号」也无「快递单号」，无法定位")
            continue

        ev = None
        if code:
            ev = db.scalar(select(ClaimEvent).where(ClaimEvent.code == code))
        if ev is None and order_no:
            ev = db.scalar(select(ClaimEvent).where(ClaimEvent.wdt_tid == order_no))
        if ev is not None:
            from .engine import merged_event_ids  # noqa: PLC0415
            if db.scalar(select(ClaimEvent.id).where(ClaimEvent.id == ev.id,
                                                    ClaimEvent.id.in_(merged_event_ids()))):
                _skip(row_meta, f"{ev.code} 已合并归档，请改在保留登记中操作")
                continue
        is_new = ev is None
        shipment = None
        if is_new or (order_no and order_no != ev.wdt_tid) or (tracking and tracking != ev.ship_logistics_no):
            if verify_shipment is None:
                _skip(row_meta, "缺少旺店通核验，未录入")
                continue
            # 旺店通核验。订单号为空时（单号只写在「快递单号」列），
            # 先用该值按「旺店通订单号」反查，失败再按「发货物流单号」反查。
            attempts = [(tracking, order_no)]
            if not order_no and tracking:
                attempts = [("", tracking), (tracking, "")]
            result = None
            reason = "未获取旺店通原始订单号"
            for lg_no, tid in attempts:
                try:
                    r = verify_shipment(lg_no, tid)
                    if r.get("ok") and r.get("shipment", {}).get("source_tid"):
                        result = r
                        break
                    reason = r.get("message") or reason
                except Exception as exc:  # noqa: BLE001
                    reason = str(exc)
            if result is None:
                _skip(row_meta, f"旺店通核验未通过：{reason}")
                continue
            try:
                shipment = result["shipment"]
                order_no = shipment["source_tid"]
                duplicate = db.scalar(select(ClaimEvent).where(ClaimEvent.wdt_tid == order_no))
                if duplicate is None and tracking:
                    duplicate = db.scalar(select(ClaimEvent).where(ClaimEvent.ship_logistics_no == tracking))
                if duplicate is not None and (ev is None or duplicate.id != ev.id):
                    raise ValueError("原始订单号已登记，请在原记录中补充")
            except Exception as exc:
                _skip(row_meta, f"旺店通核验未通过：{exc}")
                continue
        if is_new:
            if str(get("售后类型") or "").strip() not in TYPE_MAP or _money(get("赔付金额")) is None:
                _skip(row_meta, "新记录必须提供明确售后原因和仓库赔付金额")
                continue
            try:
                d = _parse_date(get("日期"))
            except ValueError as exc:
                _skip(row_meta, str(exc))
                continue
            from .engine import next_code  # noqa: PLC0415

            # 编号依据 = 订单号（与新建登记一致）；无订单号用 CLM 临时编号
            base_code = order_no
            if base_code:
                exist = db.scalar(select(ClaimEvent).where(ClaimEvent.code == base_code))
                if exist is not None:
                    _skip(row_meta, f"订单号 {order_no} 已登记（{exist.code}）")
                    continue
            ev = ClaimEvent(code=base_code or next_code(db, d), event_date=d)
            db.add(ev)

        ev.wdt_tid = order_no or ev.wdt_tid
        ev.ship_carrier = str(get("快递公司") or ev.ship_carrier or "")
        ev.ship_logistics_no = str(get("快递单号") or ev.ship_logistics_no or "")
        ev.shop_name = str(get("店铺名称") or ev.shop_name or "")
        ev.platform = platform_of(ev.shop_name) or ev.platform
        ev.customer_claim = str(get("售后内容") or ev.customer_claim or "")
        if get("备注"):
            ev.note = str(get("备注"))

        t = str(get("售后类型") or "").strip()
        if t and t in TYPE_MAP:
            ev.incident_type = TYPE_MAP[t]
        elif t:
            issues.append(f"第{i}行：售后类型「{t}」未识别，保留原值")

        ca = _money(get("赔付金额"))
        if ca is not None:
            ev.claim_amount = ca
        ra = _money(get("实际到手"))
        if ra is not None:
            ev.received_amount = ra

        duty_raw = str(get("责任方") or "").strip()
        if duty_raw in DUTY_MAP:
            ev.duty_party = DUTY_MAP[duty_raw]
            if not ev.duty_basis:
                ev.duty_basis = f"旧表责任方={duty_raw}"
        elif duty_raw:
            issues.append(f"第{i}行：责任方「{duty_raw}」未识别，置为待定")
            ev.duty_party = DutyParty.UNDECIDED

        st_raw = str(get("状态") or "").strip()
        if st_raw in STATUS_MAP:
            ev.status = STATUS_MAP[st_raw]
        elif any(x in st_raw for x in ("驳回", "无法追责")):
            ev.status = ClaimStatus.REJECTED
        elif is_new:
            ev.status = ClaimStatus.SUBMITTED

        if shipment is not None and apply_shipment is not None:
            apply_shipment(ev, shipment)
        if validate_amounts is not None:
            validate_amounts(ev)
        db.flush()

        # 多轮回复列：按「列名前缀」分组（代码不依赖任何具体主体名）。
        # 含「回复/备注」的列，去掉尾部数字后相同前缀归一组：
        # 按首次出现顺序，第1组→仓库回复，第2组→货主备注（组内列可交替出现）。
        import re as _re
        _num = _re.compile(r"\d+$")
        reply_cols = [str(h).strip() for h in raw.keys()
                      if _re.search(r"(回复|备注)", str(h or "")) and str(h or "").strip()]
        bases: list[str] = []
        for c in reply_cols:
            b = _num.sub("", c)
            if b not in bases:
                bases.append(b)
        plan: list[tuple[ReplySource, list[str]]] = []
        for i, base in enumerate(bases):
            source = ReplySource.WAREHOUSE if i == 0 else ReplySource.OWNER
            plan.append((source, [c for c in reply_cols if _num.sub("", c) == base]))
        for source, cols in plan:
            for col in cols:
                text = str(get(col) or "").strip()
                if not text or text.startswith("=DISPIMG"):
                    continue
                if any(r.source == source and r.content == text for r in ev.replies):
                    continue
                db.add(ClaimReply(event_id=ev.id, source=source, content=text))

        if is_new:
            created += 1
        else:
            updated += 1

    db.commit()
    return {
        "created": created, "updated": updated, "skipped": skipped,
        "issues": issues[:50], "issue_count": len(issues),
        "skipped_rows": skipped_rows,
    }
