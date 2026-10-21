"""REST API。"""
from __future__ import annotations

import io
import json
import logging
import uuid
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from .config import settings
from .security import (
    TOKEN_TTL_SECONDS,
    hash_password,
    make_token,
    require_admin,
    require_auth,
    verify_token,
    verify_user_password,
)
from .database import get_db
from .services.permissions import (
    Identity,
    check_fields,
    current_identity,
    identity_of,
    require_admin_identity,
    require_editor,
)
from .models import (
    ALL_FIELDS,
    ClaimEvent,
    ClaimReply,
    ClaimStatus,
    CHANNEL_LABEL,
    Channel,
    DutyChange,
    DutyParty,
    ROLE_LABEL,
    Role,
    AuditLog,
    AfterSalesType,
    Evidence,
    IncidentType,
    ReplySource,
    StockinItem,
    User,
    platform_of,
    WorkOrder,
    WorkOrderTimeline,
    WorkOrderCategory,
    WorkOrderPriority,
    WorkOrderStatus,
    WorkOrderSource,
    TimelineSource,
    WOCAT_LABEL,
    WOPRI_LABEL,
    WOSTAT_LABEL,
    TL_SRC_LABEL,
)
from .services.engine import (
    AUDIT_FIELDS,
    ENUM_LABELS,
    add_reply,
    audit_action,
    audit_diff,
    change_duty,
    find_duplicates,
    merged_event_ids,
    get_audit,
    next_code,
    OVERDUE_DAYS,
    parse_datetime,
    serialize_event,
    summarize,
)
from .services.report import (
    EXPORT_HEADERS,
    build_export_workbook,
    build_report_workbook,
    export_rows,
    parse_report_import,
)
from .services.wangdian import WangdianClient, WdtError
from .services.frontend import frontend_version, page_ready, record_ready

router = APIRouter(prefix="/api", dependencies=[Depends(require_auth)])
open_router = APIRouter(prefix="/api")

INCIDENT_ALIAS: dict[str, IncidentType] = {
    "破损": IncidentType.DAMAGE, "damage": IncidentType.DAMAGE,
    "丢件": IncidentType.LOSS, "丢失": IncidentType.LOSS, "loss": IncidentType.LOSS,
    "短少": IncidentType.SHORTAGE, "少件": IncidentType.SHORTAGE, "shortage": IncidentType.SHORTAGE,
    "错发": IncidentType.WRONG, "wrong": IncidentType.WRONG,
    "漏发/错发": IncidentType.SHORTAGE, "漏发": IncidentType.SHORTAGE,
    "未派送/派送错误": IncidentType.UNDELIVERED, "未派送": IncidentType.UNDELIVERED,
    "派送错误": IncidentType.UNDELIVERED, "undelivered": IncidentType.UNDELIVERED,
    "其他": IncidentType.OTHER, "other": IncidentType.OTHER,
}
AFTER_SALES_TYPE_ALIAS: dict[str, AfterSalesType] = {
    "原单退回": AfterSalesType.ORIGINAL_RETURN, "original_return": AfterSalesType.ORIGINAL_RETURN,
    "仅退回": AfterSalesType.RETURN_ONLY, "return_only": AfterSalesType.RETURN_ONLY,
    "退回退款": AfterSalesType.RETURN_REFUND, "return_refund": AfterSalesType.RETURN_REFUND,
}
DUTY_ALIAS: dict[str, DutyParty] = {
    "仓库": DutyParty.WAREHOUSE, "仓库方": DutyParty.WAREHOUSE, "warehouse": DutyParty.WAREHOUSE,
    "快递": DutyParty.EXPRESS, "express": DutyParty.EXPRESS,
    "共担": DutyParty.SHARED, "共担：快递+仓库": DutyParty.SHARED, "shared": DutyParty.SHARED,
    "聚小翻": DutyParty.JUXIAOFAN, "juxiaofan": DutyParty.JUXIAOFAN,
    "买家": DutyParty.BUYER, "buyer": DutyParty.BUYER,
    "人工输入": DutyParty.CUSTOM, "custom": DutyParty.CUSTOM,
    "待定": DutyParty.UNDECIDED, "undecided": DutyParty.UNDECIDED,
}
STATUS_ALIAS: dict[str, ClaimStatus] = {
    "待提交": ClaimStatus.DRAFT, "待更新数据": ClaimStatus.DRAFT, "draft": ClaimStatus.DRAFT,
    "已提交": ClaimStatus.SUBMITTED, "submitted": ClaimStatus.SUBMITTED,
    "回复中": ClaimStatus.REPLYING, "replying": ClaimStatus.REPLYING,
    "已赔付": ClaimStatus.PAID, "paid": ClaimStatus.PAID,
    "已驳回": ClaimStatus.REJECTED, "仲裁驳回": ClaimStatus.REJECTED,
    "无法追责": ClaimStatus.REJECTED, "rejected": ClaimStatus.REJECTED,
    "已放弃": ClaimStatus.GAVE_UP, "放弃": ClaimStatus.GAVE_UP, "gave_up": ClaimStatus.GAVE_UP,
}


def _dec(v: Any, field: str) -> Decimal:
    if v is None or v == "":
        return Decimal("0")
    try:
        amount = Decimal(str(v))
        if not amount.is_finite() or amount < 0:
            raise HTTPException(400, f"{field}必须是非负有限金额")
        if amount > Decimal("9999999999.99") or amount != amount.quantize(Decimal("0.01")):
            raise HTTPException(400, f"{field}最多两位小数，且不能超过金额上限")
        return amount
    except InvalidOperation as exc:
        raise HTTPException(400, f"{field} 不是合法金额：{v}") from exc


def _dec_wdt(v: Any) -> Decimal:
    """外部系统（旺店通）回传金额的归一化。

    旺店通返回的是浮点数，常带两位以上小数（如 24.900000000000002）。
    这类值不应触发「最多两位小数」的校验失败而中断整单导入，统一四舍五入到分。
    """
    try:
        amount = Decimal(str(v if v not in (None, "") else 0))
    except InvalidOperation:
        return Decimal("0")
    if not amount.is_finite() or amount < 0:
        return Decimal("0")
    if amount > Decimal("9999999999.99"):
        return Decimal("9999999999.99")
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _dt_date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    if "T" in s or " " in s:
        try:
            return datetime.fromisoformat(s).date()
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s[:10] if "T" in s else s, fmt).date()
        except ValueError:
            continue
    # 8月12日 / 2026年8月12日
    s2 = s.replace("年", "-").replace("月", "-").replace("日", "")
    parts = [p for p in s2.split("-") if p.strip()]
    try:
        if len(parts) == 2:
            m, d = int(parts[0]), int(parts[1])
            today = date.today()
            cand = date(today.year, m, d)
            return cand if cand <= today else date(today.year - 1, m, d)
        if len(parts) == 3:
            return date(int(parts[0]), int(parts[1]), int(parts[2]))
    except (ValueError, IndexError):
        pass
    raise HTTPException(400, f"日期格式无法识别：{v}")


def _dt_datetime(v: Any) -> datetime:
    if isinstance(v, datetime):
        return v.replace(second=0, microsecond=0)
    s = str(v or "").strip().replace("Z", "")
    for fmt in ("%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s[:19], fmt).replace(second=0, microsecond=0)
        except ValueError:
            continue
    raise HTTPException(400, f"登记时间格式无法识别：{v}")


def _enum(cls, v: Any, field: str, default=None):
    if v is None or v == "":
        return default
    alias = {IncidentType: INCIDENT_ALIAS, AfterSalesType: AFTER_SALES_TYPE_ALIAS,
             DutyParty: DUTY_ALIAS, ClaimStatus: STATUS_ALIAS}.get(cls, {})
    key = str(v).strip()
    if key in alias:
        return alias[key]
    try:
        return cls(key)
    except ValueError as exc:
        raise HTTPException(400, f"{field} 取值非法：{v}") from exc


# ---------------------------------------------------------------- 渠道字段权限


_ENUM_FIELD_CLS = {
    "incident_type": IncidentType,
    "after_sales_type": AfterSalesType,
    "duty_party": DutyParty,
    "status": ClaimStatus,
}
_MONEY_FIELDS = {"claim_amount", "received_amount", "refund_amount"}


def _payload_field_value(p: EventIn, name: str) -> tuple[bool, Any]:
    """把请求体里的字段解析成与数据库同型别，便于「是否真的改了」比较。

    返回 (是否可比较, 解析后的值)；解析失败时交给后续校验报 400，这里先跳过。
    """
    raw = getattr(p, name, None)
    if raw is None:
        return False, None
    try:
        if name in _ENUM_FIELD_CLS:
            return True, _enum(_ENUM_FIELD_CLS[name], raw, name)
        if name in _MONEY_FIELDS:
            return True, _dec(raw, name)
        if name == "event_date":
            return True, _dt_datetime(raw)
        if name == "closed":
            return True, bool(raw)
    except Exception:  # noqa: BLE001 — 非法值留给正式校验报错
        return False, None
    return True, str(raw).strip()


def _changed_fields(ev: ClaimEvent, p: EventIn) -> list[str]:
    """列出本次请求相对现值「确实发生变化」的字段。"""
    changed: list[str] = []
    for name in sorted(p.model_fields_set):
        if name not in EventIn.model_fields:
            continue
        ok, new = _payload_field_value(p, name)
        if not ok:
            continue
        if name == "event_date":
            if new != _dt_datetime(ev.event_time or ev.event_date):
                changed.append(name)
            continue
        old = getattr(ev, name, None)
        if name in _MONEY_FIELDS:
            old_v = _dec(old or 0, name)
            if new != old_v:
                changed.append(name)
        elif name in _ENUM_FIELD_CLS or name == "closed":
            old_v = old.value if hasattr(old, "value") else bool(old)
            if new != old_v:
                changed.append(name)
        elif str(new) != str(old or "").strip():
            changed.append(name)
    return changed


def _restrict_payload(ident: Identity, p: EventIn) -> list[str]:
    """把当前渠道不可写的字段从请求体里清空（改为「未提交」）。

    新建时用：不是自己负责的字段就当没填，避免客服顺手把仓库金额覆盖掉。
    返回被清空的字段名，供调用方按需提示。
    """
    allowed = ident.writable_fields()
    if allowed == ALL_FIELDS:
        return []
    dropped = [f for f in sorted(p.model_fields_set) if f not in allowed]
    for f in dropped:
        setattr(p, f, None)
        p.model_fields_set.discard(f)
    return dropped


# ---------------------------------------------------------------- 请求体


class EventIn(BaseModel):
    event_date: str | None = None
    wdt_tid: str | None = None
    ship_carrier: str | None = None
    ship_logistics_no: str | None = None
    shop_name: str | None = None
    refund_amount: str | None = None   # 顾客已赔（手填兜底，正常走旺店通同步回填）
    refund_reason: str | None = None
    incident_type: str | None = None
    after_sales_type: str | None = None
    customer_claim: str | None = None
    duty_party: str | None = None
    duty_party_custom: str | None = None
    duty_basis: str | None = None
    claim_amount: str | None = None
    received_amount: str | None = None
    status: str | None = None
    warehouse_ref: str | None = None
    note: str | None = None
    closed: bool | None = None


class ReplyIn(BaseModel):
    source: str
    content: str
    created_by: str = ""


class DutyIn(BaseModel):
    to_party: str
    reason: str
    duty_party_custom: str = ""


class DuplicateCorrectionIn(BaseModel):
    code: str
    wdt_tid: str
    ship_logistics_no: str


class DuplicateReviewIn(BaseModel):
    fingerprint: str
    decision: Literal["correct", "separate", "merge"]
    reason: str
    keep_code: str = ""
    corrections: list[DuplicateCorrectionIn] = []


class LoginIn(BaseModel):
    email: str
    password: str                 # 密码登录（必填）


class ChangePasswordIn(BaseModel):
    old_password: str             # 原密码（必填）
    new_password: str


class UserIn(BaseModel):
    email: str
    password: str                 # 必填：开户时由管理员设置初始密码
    is_admin: bool = False
    channel: str = ""             # warehouse / outsource / operations
    role: str = ""                # admin / editor / viewer；留空按 is_admin 推断
    display_name: str = ""


class UserPatchIn(BaseModel):
    password: str | None = None
    is_admin: bool | None = None
    channel: str | None = None
    role: str | None = None
    display_name: str | None = None
    is_active: bool | None = None


def _login_name(raw: str) -> str:
    """登录名规范化：邮箱或用户名都转小写去空白。

    用户名限制为 3-32 位字母/数字/._-，邮箱按原样（已含 @）。
    """
    name = (raw or "").strip().lower()
    if not name:
        raise HTTPException(400, "请填写登录账号")
    if "@" in name:
        if len(name) > 255 or name.startswith("@") or name.endswith("@"):
            raise HTTPException(400, "邮箱格式不合法")
        return name
    import re

    if not re.fullmatch(r"[a-z0-9][a-z0-9._-]{2,31}", name):
        raise HTTPException(
            400, "用户名需 3-32 位，只能用小写字母、数字、点、下划线或短横线（也可直接用邮箱）"
        )
    return name


def _norm_channel(raw: str | None) -> str:
    if raw is None or not str(raw).strip():
        return ""
    try:
        return Channel(str(raw).strip()).value
    except ValueError as exc:
        raise HTTPException(400, f"渠道取值非法：{raw}（仓库 / 外包客服 / 运营和客服）") from exc


def _norm_role(raw: str | None, is_admin: bool = False) -> str:
    if is_admin:
        return Role.ADMIN.value
    if raw is None or not str(raw).strip():
        return Role.EDITOR.value
    try:
        return Role(str(raw).strip()).value
    except ValueError as exc:
        raise HTTPException(400, f"角色取值非法：{raw}（管理员 / 编辑 / 只读）") from exc


# ---------------------------------------------------------------- 基础


class PageReadyIn(BaseModel):
    page_id: str
    revision: str
    columns_ready: bool = False
    shipment_form_visible: bool = False


@open_router.get("/page-version")
def get_page_version():
    return frontend_version()


@router.post("/page-ready")
def frontend_ready(body: PageReadyIn):
    if len(body.page_id) > 100 or len(body.revision) > 30:
        raise HTTPException(400, "页面版本信息格式不正确")
    record_ready(body.page_id, body.revision, body.columns_ready, body.shipment_form_visible)
    return {"ok": True}


@router.get("/page-ready/{page_id}")
def get_page_ready(page_id: str):
    return page_ready(page_id)


@open_router.get("/health")
def health(db: Session = Depends(get_db)):
    n = db.scalar(select(func.count()).select_from(ClaimEvent)) or 0
    return {"ok": True, "app": settings.app_name, "version": settings.app_version, "events": n}


@open_router.get("/meta")
def meta():
    client = WangdianClient()
    return {
        "version": settings.app_version,
        "wdt_configured": client.configured,
        "wdt_qimen_configured": client.qimen_configured,
        "incident_type": [
            {"value": x.value, "label": l} for x, l in (
                (IncidentType.DAMAGE, "破损"), (IncidentType.LOSS, "丢件"),
                (IncidentType.SHORTAGE, "漏发/错发"),
                (IncidentType.UNDELIVERED, "未派送/派送错误"))],
        "after_sales_type": [
            {"value": x.value, "label": l} for x, l in (
                (AfterSalesType.ORIGINAL_RETURN, "原单退回"),
                (AfterSalesType.RETURN_ONLY, "仅退回"),
                (AfterSalesType.RETURN_REFUND, "退回退款"))],
        "duty_party": [
            {"value": x.value, "label": l} for x, l in (
                (DutyParty.WAREHOUSE, "仓库方"), (DutyParty.EXPRESS, "快递方"),
                (DutyParty.SHARED, "共担：快递+仓库"), (DutyParty.JUXIAOFAN, "聚小翻"),
                (DutyParty.BUYER, "买家"), (DutyParty.CUSTOM, "人工输入"),
                (DutyParty.UNDECIDED, "待定"))],
        "status": [
            {"value": x.value, "label": l} for x, l in (
                (ClaimStatus.DRAFT, "待更新数据"), (ClaimStatus.SUBMITTED, "已提交仓库"),
                (ClaimStatus.REPLYING, "仓库回复中"), (ClaimStatus.PAID, "已赔付"),
                (ClaimStatus.REJECTED, "已驳回"), (ClaimStatus.GAVE_UP, "已放弃"))],
        "reply_source": [
            {"value": x.value, "label": l} for x, l in (
                (ReplySource.WAREHOUSE, "仓库回复"), (ReplySource.OWNER, "货主备注"))],
        "evidence_kind": [
            {"value": "screenshot", "label": "截图"}, {"value": "package", "label": "外箱照片"}, {"value": "inner", "label": "内物照片"},
            {"value": "track", "label": "物流轨迹"}, {"value": "chat", "label": "沟通记录"},
            {"value": "arbitration", "label": "仲裁结果"}, {"value": "stockin", "label": "入库凭证"},
            {"value": "other", "label": "其他"}],
    }


@open_router.post("/login")
def login(payload: LoginIn, db: Session = Depends(get_db)):
    """账号 + 密码登录，返回 30 天有效 token。未录入的账号一律 401。"""
    email = payload.email.strip().lower()
    user = db.scalar(select(User).where(User.email == email))
    if user is None:
        if db.scalar(select(func.count()).select_from(User)) == 0:
            raise HTTPException(
                503, "系统尚未初始化管理员账号。请在 .env 配置 ADMIN_EMAIL / ADMIN_PASSWORD 后重启。"
            )
        raise HTTPException(401, "该账号未注册，请联系管理员开通账号")
    if not user.is_active:
        raise HTTPException(401, "账号已停用，请联系管理员")
    if not payload.password:
        raise HTTPException(400, "请输入密码")
    if not user.password_hash or not verify_user_password(payload.password, user.password_hash):
        raise HTTPException(401, "密码错误")
    return _issue_token(user)


def _issue_token(user: User) -> dict:
    return {"ok": True, "token": make_token(user.email), "expires_in": TOKEN_TTL_SECONDS}


@router.get("/me")
def me(email: str = Depends(require_auth), db: Session = Depends(get_db)):
    if settings.disable_auth:
        # 开发期免登录：直接返回虚拟管理员身份，并标记 dev_mode 供前端识别。
        ident = Identity(email=email, role=Role.ADMIN.value, dev_mode=True)
        data = ident.to_dict()
        data.update({"ok": True, "has_password": False})
        return data
    user = db.scalar(select(User).where(User.email == email))
    if not user:
        return {"ok": False}
    ident = identity_of(user, email)
    data = ident.to_dict()
    data.update({"ok": True, "has_password": bool(user.password_hash),
                 "display_name": user.display_name or ""})
    return data


@router.post("/auth/change-password")
def change_password(payload: ChangePasswordIn, email: str = Depends(require_auth), db: Session = Depends(get_db)):
    """登录后修改密码，需校验原密码。"""
    user = db.scalar(select(User).where(User.email == email))
    if user is None:
        raise HTTPException(401, "未登录")
    if not user.password_hash or not verify_user_password(payload.old_password, user.password_hash):
        raise HTTPException(400, "原密码错误")
    if len(payload.new_password) < 6:
        raise HTTPException(400, "新密码至少 6 位")
    user.password_hash = hash_password(payload.new_password)
    db.commit()
    return {"ok": True, "message": "密码已更新"}


@router.get("/users")
def list_users(admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    users = list(db.scalars(select(User).order_by(User.id)).all())
    return {
        "rows": [
            {"id": u.id, "email": u.email, "is_admin": bool(u.is_admin),
             "is_active": bool(u.is_active), "has_password": bool(u.password_hash),
             "channel": u.channel or "", "role": u.role or ("admin" if u.is_admin else "editor"),
             "display_name": u.display_name or "",
             "created_at": u.created_at.isoformat() if u.created_at else ""}
            for u in users
        ],
        "channels": [{"value": c.value, "label": CHANNEL_LABEL[c.value]} for c in Channel],
        "roles": [{"value": r.value, "label": ROLE_LABEL[r.value]} for r in Role],
    }


@router.post("/users")
def create_user(payload: UserIn, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    email = _login_name(payload.email)
    if len(payload.password) < 6:
        raise HTTPException(400, "初始密码至少 6 位")
    if db.scalar(select(User).where(User.email == email)):
        raise HTTPException(409, "该账号已存在")
    channel = _norm_channel(payload.channel)
    role = _norm_role(payload.role, payload.is_admin)
    db.add(User(email=email, password_hash=hash_password(payload.password),
                is_admin=(role == Role.ADMIN.value), channel=channel, role=role,
                display_name=(payload.display_name or "").strip()[:64]))
    db.commit()
    return {"ok": True, "email": email, "has_password": True, "channel": channel, "role": role}


@router.patch("/users/{uid}")
def update_user(uid: int, payload: UserPatchIn,
                admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    """修改账号的渠道 / 角色 / 密码 / 启停。"""
    user = db.get(User, uid)
    if user is None:
        raise HTTPException(404, "账号不存在")
    # 先算出目标状态再校验，校验不过时不动 ORM 对象，避免半改状态。
    if payload.role is not None:
        # role 为准：管理员降级/升级都走这里，is_admin 只作历史兼容。
        new_role = _norm_role(payload.role, bool(payload.is_admin))
    elif payload.is_admin is not None:
        new_role = Role.ADMIN.value if payload.is_admin else (
            user.role if user.role != Role.ADMIN.value else Role.EDITOR.value)
    else:
        new_role = user.role
    new_channel = _norm_channel(payload.channel) if payload.channel is not None else user.channel
    new_active = user.is_active if payload.is_active is None else payload.is_active
    if user.id == admin.id and not new_active:
        raise HTTPException(400, "不能停用自己")
    if payload.password is not None and len(payload.password) < 6:
        raise HTTPException(400, "密码至少 6 位")
    # 至少保留一个可用管理员
    if new_role != Role.ADMIN.value or not new_active:
        others = [u for u in db.scalars(select(User).where(User.is_admin, User.is_active)).all()
                  if u.id != user.id]
        if not others:
            raise HTTPException(400, "至少保留一个启用的管理员账号")

    user.channel = new_channel
    user.role = new_role
    user.is_admin = new_role == Role.ADMIN.value
    if payload.display_name is not None:
        user.display_name = payload.display_name.strip()[:64]
    if payload.password is not None:
        user.password_hash = hash_password(payload.password)
    user.is_active = new_active
    db.commit()
    return {"ok": True, "id": user.id, "channel": user.channel, "role": user.role,
            "is_admin": bool(user.is_admin), "is_active": bool(user.is_active),
            "display_name": user.display_name}


@router.delete("/users/{uid}")
def delete_user(uid: int, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    user = db.get(User, uid)
    if user is None:
        raise HTTPException(404, "账号不存在")
    if user.id == admin.id:
        raise HTTPException(400, "不能删除自己")
    admins = list(db.scalars(select(User).where(User.is_admin, User.is_active)).all())
    if user.is_admin and len(admins) <= 1:
        raise HTTPException(400, "至少保留一个管理员账号")
    db.delete(user)
    db.commit()
    return {"ok": True}


def _apply_filters(q, month: str, start: str, end: str):
    q = q.where(~ClaimEvent.id.in_(merged_event_ids()))
    if month:
        try:
            y, m = month.split("-")
            y = int(y)
            # Use EXTRACT so month filtering works on both SQLite and PostgreSQL.
            q = q.where(
                (func.extract("year", ClaimEvent.event_date) == y)
                & (func.extract("month", ClaimEvent.event_date) == int(m))
            )
        except ValueError as exc:
            raise HTTPException(400, "month 格式应为 2026-08") from exc
    if start:
        q = q.where(ClaimEvent.event_date >= _dt_date(start))
    if end:
        q = q.where(ClaimEvent.event_date < _dt_date(end) + timedelta(days=1))
    return q


@router.get("/events")
def list_events(
    db: Session = Depends(get_db),
    month: str = Query(default=""),
    start: str = Query(default=""),
    end: str = Query(default=""),
    shop: str = Query(default=""),
    carrier: str = Query(default=""),
    keyword: str = Query(default=""),
    incident_type: str = Query(default=""),
    duty_party: str = Query(default=""),
    status: str = Query(default=""),
    only_pending_duty: bool = Query(default=False),
    only_gap: bool = Query(default=False),
    only_overdue: bool = Query(default=False),
    limit: int = Query(default=100, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
):
    q = _apply_filters(select(ClaimEvent), month, start, end)
    if shop:
        q = q.where(ClaimEvent.shop_name.like(f"%{shop}%"))
    if carrier:
        q = q.where(ClaimEvent.ship_carrier.like(f"%{carrier}%"))
    if incident_type:
        q = q.where(ClaimEvent.incident_type == _enum(IncidentType, incident_type, "类型"))
    if duty_party:
        q = q.where(ClaimEvent.duty_party == _enum(DutyParty, duty_party, "责任方"))
    if status:
        q = q.where(ClaimEvent.status == _enum(ClaimStatus, status, "状态"))
    if keyword:
        like = f"%{keyword}%"
        q = q.where(
            ClaimEvent.wdt_tid.like(like) | ClaimEvent.ship_logistics_no.like(like)
            | ClaimEvent.code.like(like) | ClaimEvent.customer_claim.like(like)
        )
    if only_overdue:
        q = q.where(ClaimEvent.status.in_([ClaimStatus.SUBMITTED, ClaimStatus.REPLYING]))
        q = q.where(ClaimEvent.event_date <= (date.today() - timedelta(days=OVERDUE_DAYS)))
    evs_all = list(db.scalars(q.order_by(ClaimEvent.event_date.desc(), ClaimEvent.event_time.desc().nullslast(), ClaimEvent.id.desc())).unique().all())
    rows = [serialize_event(e) for e in evs_all]
    if only_pending_duty:
        rows = [r for r in rows if r["duty_party"] == "undecided"]
    if only_gap:
        rows = [r for r in rows if Decimal(r["gap_amount"]) > 0]
    total = len(rows)
    rows = rows[offset:offset + limit]
    return {"total": total, "rows": rows}


@router.get("/page-version")
def page_version():
    return frontend_version()

@router.get("/stats")
def stats(
    db: Session = Depends(get_db),
    month: str = Query(default=""),
    start: str = Query(default=""),
    end: str = Query(default=""),
):
    q = _apply_filters(select(ClaimEvent), month, start, end)
    evs = list(db.scalars(q).unique().all())
    data = summarize(evs)
    data["duplicates"] = find_duplicates(db)
    return data


# ---------------------------------------------------------------- 事件 CRUD


@router.get("/duplicates")
def duplicate_groups(db: Session = Depends(get_db)):
    return {"groups": find_duplicates(db)}


@router.post("/duplicates/review")
def review_duplicate(payload: DuplicateReviewIn, ident: Identity = Depends(require_editor),
                     db: Session = Depends(get_db)):
    actor = ident.email
    check_fields(ident, ["wdt_tid", "ship_logistics_no"])
    reason = payload.reason.strip()
    if len(reason) < 6:
        raise HTTPException(400, "请填写至少6个字的核对依据，说明修正内容或分别申报的原因")
    group = next((g for g in find_duplicates(db)
                  if g["fingerprint"] == payload.fingerprint), None)
    if group is None:
        raise HTTPException(409, "记录已变化或已处理，请刷新重复核对窗口")
    # Serialize edits/reviews of these records across concurrent operators.
    events = list(db.scalars(select(ClaimEvent).where(
        ClaimEvent.code.in_(group["codes"])).order_by(ClaimEvent.id).with_for_update()).all())
    for ev in events:
        db.refresh(ev)
    if not any(g["fingerprint"] == payload.fingerprint for g in find_duplicates(db)):
        raise HTTPException(409, "记录已变化或已处理，请重新核对")
    if payload.decision == "merge":
        keep = next((ev for ev in events if ev.code == payload.keep_code), None)
        if keep is None:
            raise HTTPException(400, "请选择本组中要保留的登记")
        for ev in events:
            if ev.id == keep.id:
                continue
            # Preserve original data, attachments, refunds and audit history in place.
            # Duplicate monetary amounts must never be automatically added together.
            audit_action(db, ev, actor, "duplicate_merge", field="合并归档",
                         old_value=str(keep.id), new_value=json.dumps(
                             {"keep_code": keep.code, "reason": reason}, ensure_ascii=False))
        audit_action(db, keep, actor, "duplicate_merge", field="合并保留",
                     old_value=payload.fingerprint, new_value=json.dumps({
                         "source_codes": [ev.code for ev in events if ev.id != keep.id],
                         "reason": reason, "amount_policy": "保留所选记录的金额，不累加重复金额"
                     }, ensure_ascii=False))
    elif payload.decision == "separate":
        for ev in events:
            audit_action(db, ev, actor, "duplicate_review", field="确认独立申报",
                         old_value=payload.fingerprint,
                         new_value=json.dumps({"field": group["field"], "value": group["value"],
                                               "codes": group["codes"], "reason": reason}, ensure_ascii=False))
    else:
        corrections = {row.code: row for row in payload.corrections}
        if set(corrections) != set(group["codes"]) or len(corrections) != len(payload.corrections):
            raise HTTPException(400, "必须提交该组全部记录的核对结果，不可修改组外记录")
        changed = False
        for ev in events:
            row = corrections[ev.code]
            tid, lno = row.wdt_tid.strip(), row.ship_logistics_no.strip()
            if not tid and not lno:
                raise HTTPException(400, "每条记录至少保留订单号或发货单号，不能清空号码跳过核对")
            if len(tid) > 64 or len(lno) > 64:
                raise HTTPException(400, "订单号和发货单号不能超过64个字符")
            verified = _lookup_shipment(lno, tid)
            if not verified.get("ok") or not verified.get("shipment", {}).get("source_tid"):
                raise HTTPException(409, "修正后的单号未通过旺店通核验，请核对")
            shipment = verified["shipment"]
            if tid and shipment["source_tid"] != tid:
                raise HTTPException(409, "修正后的原始订单号与旺店通结果不一致")
            tid = shipment["source_tid"]
            lno = shipment.get("logistics_no") or lno
            before = {a: getattr(ev, a) for a, _ in AUDIT_FIELDS}
            changed |= ev.wdt_tid != tid or ev.ship_logistics_no != lno
            ev.wdt_tid, ev.ship_logistics_no = tid, lno
            audit_diff(db, ev, actor, before)
            audit_action(db, ev, actor, "duplicate_review", field="修正关联单号",
                         old_value=payload.fingerprint, new_value=reason)
        if not changed:
            raise HTTPException(400, "尚未修正任何号码；如确为不同事故，请填写依据后确认独立申报")
    db.commit()
    return {"ok": True, "groups": find_duplicates(db)}


@router.get("/events/{code}")
def get_event(code: str, db: Session = Depends(get_db)):
    ev = _get(db, code, include_merged=True)
    result = serialize_event(ev)
    marker = db.scalar(select(AuditLog).where(AuditLog.event_id == ev.id,
        AuditLog.action == "duplicate_merge", AuditLog.field == "合并归档"))
    if marker:
        target = db.get(ClaimEvent, int(marker.old_value))
        result["merged_into"] = {"code": target.code} if target else {"code": "已归档"}
    sources = db.scalars(select(ClaimEvent).where(ClaimEvent.id.in_(
        select(AuditLog.event_id).where(AuditLog.action == "duplicate_merge",
            AuditLog.field == "合并归档", AuditLog.old_value == str(ev.id))
    )).order_by(ClaimEvent.id)).all()
    result["merged_sources"] = [{"code": src.code, "event_date": str(src.event_date),
        "incident_type": src.incident_type.value, "claim_amount": f"{src.claim_amount:.2f}",
        "received_amount": f"{src.received_amount:.2f}"} for src in sources]
    return result


@router.get("/events/{code}/audit")
def event_audit(code: str, db: Session = Depends(get_db)):
    """单事件操作审计记录（按时间倒序）。"""
    return {"rows": get_audit(db, code)}


@router.get("/audit")
def global_audit(db: Session = Depends(get_db), limit: int = Query(default=200)):
    """全局最近操作审计（默认 200 条）。"""
    rows = db.scalars(
        select(AuditLog).order_by(AuditLog.id.desc()).limit(limit)
    ).all()
    return {"rows": [r.as_dict() for r in rows]}


@router.post("/events")
def create_event(payload: EventIn, ident: Identity = Depends(require_editor), db: Session = Depends(get_db)):
    actor = ident.email
    if not ident.can_create:
        raise HTTPException(
            403,
            "登记新售后由「运营和客服 / 外包客服」负责；仓库账号请打开已有登记补充入库与仓库信息",
        )
    # 渠道字段过滤：不是自己负责的字段当作没填，避免顺手覆盖别人口径。
    _restrict_payload(ident, payload)
    if not (payload.incident_type or "").strip():
        raise HTTPException(400, "请主动选择售后类型，不能默认按破损登记")
    if not (payload.event_date or "").strip():
        raise HTTPException(400, "请填写登记日期（售后时间）")
    amount_missing = False
    if payload.claim_amount is None or not str(payload.claim_amount).strip():
        if ident.can_edit_field("claim_amount"):
            raise HTTPException(400, "请填写仓库赔付金额")
        # 登记渠道不负责金额（如外包客服）：先按 0 建单并标记待补，由运营补录。
        payload.claim_amount = "0"
        amount_missing = True
    # 新登记必须先由旺店通核验，不能仅凭手填号码形成无依据台账。
    lookup_key = (payload.wdt_tid or "").strip() or (payload.ship_logistics_no or "").strip()
    if not lookup_key:
        raise HTTPException(400, "请先填写发货物流单号或原始订单号，并完成旺店通核验")
    verified = _lookup_shipment(payload.ship_logistics_no or "", payload.wdt_tid or "")
    if not verified.get("ok"):
        raise HTTPException(409, "旺店通尚未查到对应的发货/订单依据，暂不能登记；请核对号码后等待查询完成")
    shipment = verified["shipment"]
    if not shipment.get("source_tid"):
        raise HTTPException(409, "旺店通未返回原始订单号，暂不能登记；请核对发货单号或等待旺店通数据完整返回")
    # 登记编号唯一依据改为旺店通返回的原始订单号，不接受手工号码覆盖。
    payload.wdt_tid = shipment["source_tid"]
    if not (payload.ship_logistics_no or "").strip() and shipment.get("logistics_no"):
        payload.ship_logistics_no = shipment["logistics_no"]
    if not (payload.ship_carrier or "").strip() and shipment.get("carrier"):
        payload.ship_carrier = shipment["carrier"]
    if not (payload.shop_name or "").strip() and shipment.get("shop_name"):
        payload.shop_name = shipment["shop_name"]
    duplicate_filters = []
    if (payload.wdt_tid or "").strip():
        duplicate_filters.append(func.trim(ClaimEvent.wdt_tid) == payload.wdt_tid.strip())
    if (payload.ship_logistics_no or "").strip():
        duplicate_filters.append(func.trim(ClaimEvent.ship_logistics_no) == payload.ship_logistics_no.strip())
    if duplicate_filters:
        existing = db.scalars(select(ClaimEvent).where(
            or_(*duplicate_filters), ~ClaimEvent.id.in_(merged_event_ids())
        ).order_by(ClaimEvent.id).limit(1)).first()
        if existing is not None:
            raise HTTPException(
                409,
                f"发现重复登记：已有记录 {existing.code} 使用相同订单号或发货单号，请先核对后再登记",
            )
    d = _dt_date(payload.event_date)
    event_time = _dt_datetime(payload.event_date)
    # 编号依据 = 旺店通订单号（用户确认）；未填订单号的线下登记用 CLM 临时编号兜底
    base_code = (payload.wdt_tid or "").strip()
    if base_code:
        dup = db.scalar(select(ClaimEvent).where(ClaimEvent.code == base_code))
        if dup is not None:
            raise HTTPException(
                409,
                "订单号 " + base_code + " 已登记过（" + dup.code + "）。"
                "同一订单的多笔退款应合并在同一张登记里，请打开原登记补充，勿重复新建",
            )
        code = base_code
    else:
        code = next_code(db, d)
    ev = ClaimEvent(code=code, event_date=d, event_time=event_time,
                    order_amount=_dec_wdt(shipment.get("goods_total_amount") or shipment.get("goods_amount") or 0))
    # 登记人维度：谁建的单、属于哪个渠道，报表按此拆分（历史记录为空→展示「历史导入」）。
    ev.created_by = ident.email
    ev.created_channel = ident.channel
    _fill(ev, payload, creating=True)
    _apply_shipment(ev, shipment)
    if amount_missing:
        # 金额不由本渠道负责：标记待补，报表汇总只统计已知金额。
        try:
            snap = json.loads(ev.shipment_snapshot or "{}")
        except (ValueError, TypeError):
            snap = {}
        mi = snap.setdefault("_manual_import", {})
        mi["amount_missing"] = True
        ev.shipment_snapshot = json.dumps(snap, ensure_ascii=False)
    if payload.duty_party:
        ev.duty_party = _enum(DutyParty, payload.duty_party, "责任方")
        if ev.duty_party == DutyParty.CUSTOM:
            ev.duty_party_custom = (payload.duty_party_custom or "").strip()
            if not ev.duty_party_custom:
                raise HTTPException(400, "选择人工输入时请填写责任方")
        ev.duty_basis = payload.duty_basis or ""
    db.add(ev)
    db.commit()
    db.refresh(ev)
    audit_action(db, ev, actor, "create", field="-", new_value=f"新建 {ev.code}")
    db.commit()
    return {"ok": True, "event": serialize_event(ev)}


@router.put("/events/{code}")
def update_event(code: str, payload: EventIn, ident: Identity = Depends(require_editor), db: Session = Depends(get_db)):
    actor = ident.email
    ev = _get(db, code)
    # 渠道字段权限：真正改动了不属于本渠道的字段才拦，未变的原样回传不报错。
    changed = _changed_fields(ev, payload)
    if changed:
        check_fields(ident, changed)
    _restrict_payload(ident, payload)
    if payload.wdt_tid is not None and not payload.wdt_tid.strip():
        raise HTTPException(400, "原始订单号不能为空，必须由旺店通核验")
    new_tid = payload.wdt_tid.strip() if payload.wdt_tid is not None else ev.wdt_tid
    new_lno = payload.ship_logistics_no.strip() if payload.ship_logistics_no is not None else ev.ship_logistics_no
    shipment = None
    if new_tid != ev.wdt_tid or new_lno != ev.ship_logistics_no:
        verified = _lookup_shipment(new_lno or "", new_tid or "")
        if not verified.get("ok"):
            raise HTTPException(409, "修改后的单号尚未通过旺店通核验")
        shipment = verified["shipment"]
        if not shipment.get("source_tid") or (new_tid and shipment["source_tid"] != new_tid):
            raise HTTPException(409, "原始订单号与旺店通返回结果不一致")
        payload.wdt_tid = shipment["source_tid"]
        payload.ship_logistics_no = shipment.get("logistics_no") or new_lno
        keys = [func.trim(ClaimEvent.wdt_tid) == payload.wdt_tid]
        if payload.ship_logistics_no:
            keys.append(func.trim(ClaimEvent.ship_logistics_no) == payload.ship_logistics_no)
        if db.scalar(select(ClaimEvent.id).where(ClaimEvent.id != ev.id,
                ~ClaimEvent.id.in_(merged_event_ids()), or_(*keys)).limit(1)):
            raise HTTPException(409, "修改后的订单号或发货单号已登记，请先核对重复记录")
    before = {a: getattr(ev, a) for a, _ in AUDIT_FIELDS}
    _fill(ev, payload)
    if shipment:
        _apply_shipment(ev, shipment)
    try:
        _maybe_change_duty(db, ev, payload)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    db.commit()
    db.refresh(ev)
    audit_diff(db, ev, actor, before)
    db.commit()
    return {"ok": True, "event": serialize_event(ev)}


def _get(db: Session, code: str, *, include_merged: bool = False) -> ClaimEvent:
    ev = db.scalar(select(ClaimEvent).where(ClaimEvent.code == code))
    if ev is None:
        raise HTTPException(404, f"事件不存在：{code}")
    if not include_merged and db.scalar(select(AuditLog.id).where(
        AuditLog.event_id == ev.id, AuditLog.action == "duplicate_merge", AuditLog.field == "合并归档")):
        raise HTTPException(409, "此登记已合并归档，请在保留的登记中继续操作")
    return ev


def _fill(ev: ClaimEvent, p: EventIn, creating: bool = False) -> None:
    if "incident_type" in p.model_fields_set and not (p.incident_type or "").strip():
        raise HTTPException(400, "请选择售后类型")
    if p.event_date:
        ev.event_date = _dt_date(p.event_date)
        ev.event_time = _dt_datetime(p.event_date)
    for f in ("wdt_tid", "ship_carrier", "ship_logistics_no", "shop_name",
              "customer_claim", "duty_basis", "warehouse_ref", "note", "refund_reason"):
        v = getattr(p, f)
        if v is not None:
            setattr(ev, f, v)
    if p.refund_amount is not None:
        ev.refund_amount = _dec(p.refund_amount, "顾客已赔")
    if p.shop_name:
        ev.platform = platform_of(p.shop_name)
    if p.incident_type:
        ev.incident_type = _enum(IncidentType, p.incident_type, "类型")
    if p.after_sales_type is not None:
        ev.after_sales_type = _enum(AfterSalesType, p.after_sales_type, "售后类型") if p.after_sales_type else None
    if p.duty_party is not None and creating:
        new_party = _enum(DutyParty, p.duty_party, "责任方")
        custom = (p.duty_party_custom or "").strip() if new_party == DutyParty.CUSTOM else ""
        if new_party == DutyParty.CUSTOM and not custom:
            raise HTTPException(400, "选择人工输入时请填写责任方")
        ev.duty_party = new_party
        ev.duty_party_custom = custom
    if p.claim_amount is not None:
        if not str(p.claim_amount).strip():
            raise HTTPException(400, "仓库赔付金额不能为空，请填写金额")
        ev.claim_amount = _dec(p.claim_amount, "仓库赔付金额")
    if p.received_amount is not None:
        ev.received_amount = _dec(p.received_amount, "实收金额")
        # 实收不能超过申请（口径校验）
        if ev.received_amount > ev.claim_amount:
            raise HTTPException(400, f"实收({ev.received_amount}) 大于申请({ev.claim_amount})，请检查")
    if p.status:
        st = _enum(ClaimStatus, p.status, "状态")
        ev.status = st
    _validate_amounts(ev)
    if p.closed is not None:
        ev.closed = p.closed
    return ev


def _validate_amounts(ev: ClaimEvent) -> None:
    claim = _dec(ev.claim_amount or 0, "仓库赔付金额")
    received = _dec(ev.received_amount or 0, "实收金额")
    _dec(ev.refund_amount or 0, "顾客已赔")
    if received > claim:
        raise HTTPException(400, "实收金额不能大于仓库赔付金额，请核对金额")
    if ev.status == ClaimStatus.PAID and (received <= 0 or received < claim):
        raise HTTPException(400, "已赔付状态需要实际收款并足额，请先核对实收金额")


def _maybe_change_duty(db: Session, ev: ClaimEvent, p: EventIn) -> None:
    """PUT 路径的责任改判：只在新旧不一致且填了依据时走留痕变更。"""
    if p.duty_party is None:
        return
    new_party = _enum(DutyParty, p.duty_party, "责任方")
    custom = (p.duty_party_custom or "").strip() if new_party == DutyParty.CUSTOM else ""
    if new_party == DutyParty.CUSTOM and not custom:
        raise HTTPException(400, "选择人工输入时请填写责任方")
    if new_party == ev.duty_party and custom == (ev.duty_party_custom or ""):
        return
    if not (p.duty_basis or "").strip():
        raise ValueError("修改责任方必须填写判定依据")
    change_duty(db, ev, new_party, p.duty_basis or "")
    ev.duty_party_custom = custom


@router.delete("/events/{code}")
def delete_event(code: str, ident: Identity = Depends(require_admin_identity), db: Session = Depends(get_db)):
    actor = ident.email
    ev = _get(db, code)
    if db.scalar(select(AuditLog.id).where(AuditLog.action == "duplicate_merge",
        AuditLog.field == "合并归档", AuditLog.old_value == str(ev.id))):
        raise HTTPException(409, "此登记关联已合并的原始记录，请保留以便追溯")
    audit_action(db, ev, actor, "delete", field="-", new_value=f"删除 {ev.code}")
    db.delete(ev)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- 回复流


@router.post("/events/{code}/replies")
def post_reply(code: str, payload: ReplyIn, ident: Identity = Depends(require_editor), db: Session = Depends(get_db)):
    actor = ident.email
    ev = _get(db, code)
    if not payload.content.strip():
        raise HTTPException(400, "回复内容不能为空")
    raw = str(payload.source).strip()
    if raw in ("仓库", "仓库回复", "warehouse"):
        src = ReplySource.WAREHOUSE
    elif raw in ("自己", "货主", "货主备注", "自己备注", "owner", "备注"):
        src = ReplySource.OWNER
    else:
        raise HTTPException(400, f"回复来源取值非法：{payload.source}（仓库 / 自己）")
    # 渠道约束：仓库账号只能记录仓库方回复；外包客服只能记货主备注；
    # 运营和客服与管理员可两边都记（负责把群里的仓库反馈抄进系统）。
    if not ident.is_admin and ident.channel == "warehouse" and src != ReplySource.WAREHOUSE:
        raise HTTPException(403, "仓库账号只能记录「仓库回复」，货主备注请由客服渠道填写")
    if not ident.is_admin and ident.channel == "outsource" and src != ReplySource.OWNER:
        raise HTTPException(403, "外包客服渠道只能记录「货主备注」，仓库方回复请由仓库账号填写")
    add_reply(db, ev, source=src, content=payload.content,
              created_by=(payload.created_by or ident.display_name or ident.email))
    db.commit()
    db.refresh(ev)
    audit_action(db, ev, actor, "reply", field="-", new_value=payload.content[:80])
    db.commit()
    return {"ok": True, "event": serialize_event(ev)}


@router.post("/events/{code}/duty")
def change_duty_api(code: str, payload: DutyIn, ident: Identity = Depends(require_editor), db: Session = Depends(get_db)):
    actor = ident.email
    check_fields(ident, ["duty_party", "duty_party_custom", "duty_basis"])
    ev = _get(db, code)
    old_party = ev.duty_party
    try:
        to_party = _enum(DutyParty, payload.to_party, "责任方")
        custom = payload.duty_party_custom.strip() if to_party == DutyParty.CUSTOM else ""
        if to_party == DutyParty.CUSTOM and not custom:
            raise ValueError("选择人工输入时请填写责任方")
        change_duty(db, ev, to_party, payload.reason)
        ev.duty_party_custom = custom
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    db.commit()
    db.refresh(ev)
    audit_action(db, ev, actor, "duty", field="责任方",
                 old_value=ENUM_LABELS["duty_party"].get(old_party.value, old_party.value),
                 new_value=ENUM_LABELS["duty_party"].get(ev.duty_party.value, ev.duty_party.value))
    db.commit()
    return {"ok": True, "event": serialize_event(ev)}


# ---------------------------------------------------------------- 旺店通同步


def _return_detection_from_refunds(client: WangdianClient, rows: list[dict[str, Any]]) -> dict[str, str]:
    """登记前的轻量退回识别：只读旺店通退款快照，不写库。"""
    if not rows:
        return {"state": "none", "label": "", "evidence": "", "suggested_type": ""}
    parsed = [client.parse_refund(row) for row in rows]
    if any((x["return_logistics_no"] or "").strip() or x["return_goods_count"] > 0 for x in parsed):
        suggested = "return_refund" if any(x["refund_amount"] > 0 for x in parsed) else "return_only"
        return {"state": "shipping", "label": "退货中·待入库", "evidence": "退款单含退货物流或退回数量；原单退回需要人工核实", "suggested_type": suggested}
    text = " ".join(f"{x['refund_reason']} {x['refund_status']}" for x in parsed)
    if any(k in text for k in ("退货", "退回", "换货", "拒收")):
        return {"state": "suspected", "label": "疑似退货·待核实", "evidence": "退款原因/状态含退货关键词", "suggested_type": "return_refund"}
    return {"state": "none", "label": "", "evidence": "", "suggested_type": ""}


def _lookup_shipment(logistics_no: str = "", src_tid: str = "") -> dict[str, Any]:
    client = WangdianClient()
    logistics_no, src_tid = logistics_no.strip(), src_tid.strip()
    if not logistics_no and not src_tid:
        raise HTTPException(400, "请填写发货物流单号或原始订单号")
    try:
        rows = client.find_sales_stockouts(src_tid=src_tid, logistics_no=logistics_no)
    except (WdtError, ValueError) as exc:
        raise HTTPException(502, str(exc)) from exc
    if not rows:
        message = "旺店通接口未返回与该发货物流号/原始订单号匹配的数据。"
        if not client.qimen_configured:
            message += "当前使用普通接口；淘宝、天猫订单需开通并配置旺店通奇门自定义接口。"
        return {"ok": False, "stage": "none", "message": message,
                "channel": "qimen" if client.qimen_configured else "standard"}
    shipments = [client.parse_sales_stockout(row) for row in rows]
    if src_tid:
        shipments = [s for s in shipments if src_tid in [x.strip() for x in s.get("source_tid", "").replace(";", ",").split(",")]]
    if not shipments:
        return {"ok": False, "stage": "none", "message": "原始订单号与发货物流单号不一致，请核对"}
    source_ids = {s.get("source_tid", "") for s in shipments}
    if len(source_ids) > 1 or (not src_tid and any("," in s or ";" in s for s in source_ids)):
        return {"ok": False, "stage": "ambiguous", "message": "查到多个原始订单，请同时填写原始订单号核对，暂不能登记"}
    shipment = shipments[0]
    if src_tid:
        shipment["source_tid"] = src_tid
    if not (shipment.get("source_tid") or "").strip():
        return {"ok": False, "stage": "none", "message": "旺店通未返回原始订单号，暂不能登记，请等待查询或核对号码"}
    shipment.pop("raw", None)
    shipment["platform"] = platform_of(shipment["shop_name"])
    shipment["return_detection"] = {"state": "none", "label": "", "evidence": "", "suggested_type": ""}
    # 登记弹窗需要即时给出“原单退回”关键提示；退款接口异常不影响发货核验本身。
    if shipment.get("source_tid"):
        try:
            refund_rows = client.find_refunds(tid=shipment["source_tid"])
            shipment["return_detection"] = _return_detection_from_refunds(client, refund_rows)
        except (WdtError, ValueError) as exc:
            logging.getLogger("claims").warning("登记前退回识别失败 tid=%s: %s", shipment["source_tid"], exc)
    return {"ok": True, "stage": "shipment", "shipment": shipment,
            "match_count": len(rows), "channel": "qimen" if client.qimen_configured else "standard",
            "message": "已从旺店通带出订单和发货信息"}


def _apply_shipment(ev: ClaimEvent, shipment: dict[str, Any]) -> None:
    for field, key in (("wdt_tid", "source_tid"), ("shop_name", "shop_name"),
                       ("ship_carrier", "carrier"), ("ship_logistics_no", "logistics_no")):
        if not (getattr(ev, field) or "").strip() and shipment.get(key):
            setattr(ev, field, shipment[key])
    if ev.shop_name:
        ev.platform = platform_of(ev.shop_name) or ev.platform
    if "goods_total_amount" in shipment:
        ev.order_amount = _dec_wdt(shipment.get("goods_total_amount") or 0)
    snapshot = {key: value for key, value in shipment.items() if key != "raw"}
    ev.shipment_snapshot = json.dumps(snapshot, ensure_ascii=False, default=str)


@router.get("/wdt/shipment")
def lookup_shipment(logistics_no: str = "", src_tid: str = ""):
    """独立查询发货信息，不要求退款，也不创建理赔记录。"""
    return _lookup_shipment(logistics_no, src_tid)


@router.post("/events/{code}/sync-shipment")
def sync_shipment(code: str, ident: Identity = Depends(require_editor), db: Session = Depends(get_db)):
    actor = ident.email
    ev = _get(db, code)
    result = _lookup_shipment(ev.ship_logistics_no or "", ev.wdt_tid or "")
    if result["ok"]:
        _apply_shipment(ev, result["shipment"])
        ev.wdt_synced_at = datetime.now()
        audit_action(db, ev, actor, "sync", field="-",
                     new_value=f"查询发货信息：运单 {ev.ship_logistics_no}，原始订单 {ev.wdt_tid}")
        db.commit()
        db.refresh(ev)
    result["event"] = serialize_event(ev)
    return result


def _persist_refunds(db: Session, ev: ClaimEvent, client: WangdianClient, rows: list,
                     *, only_fill_empty: bool = False) -> Decimal:
    """退款单落库（RefundRecord / RefundItem）+ 回填事件主快照。

    退货物流/退货单号/退货件数都来自退款单，所以「补退货信息」本质就是「拉退款单」。
    only_fill_empty=True 时只补空值，不覆盖人工填写的金额/原因（批量补取用）。
    调用方负责 commit。
    """
    from .models import RefundItem, RefundRecord  # noqa: PLC0415

    total_amount = Decimal("0")
    for row in rows:
        pr = client.parse_refund(row)
        rec = db.scalar(
            select(RefundRecord).where(RefundRecord.refund_no == pr["wdt_refund_no"])
        )
        if rec is None:
            rec = RefundRecord(event_id=ev.id, refund_no=pr["wdt_refund_no"])
            db.add(rec)
        rec.refund_amount = pr["refund_amount"]
        rec.refund_reason = pr["refund_reason"]
        rec.refund_status = pr["refund_status"]
        rec.return_carrier = pr["return_carrier"]
        rec.return_logistics_no = pr["return_logistics_no"]
        rec.return_goods_count = pr["return_goods_count"]
        rec.refund_time = pr["refund_time"]
        rec.raw_json = json.dumps(row, ensure_ascii=False, default=str)
        if pr.get("refund_status") == "5":
            total_amount += pr["refund_amount"]

    # 主快照取第一笔（兼容旧展示字段），顾客已赔 = 全部退款单合计
    parsed = client.parse_refund(rows[0])
    ev.wdt_refund_no = parsed["wdt_refund_no"]
    if parsed["shop_name"]:
        ev.shop_name = parsed["shop_name"]
        ev.platform = platform_of(parsed["shop_name"]) or ev.platform
    if total_amount > 0 and not (only_fill_empty and (ev.refund_amount or 0) > 0):
        if ev.order_amount and ev.order_amount > 0:
            ev.refund_amount = min(total_amount, ev.order_amount)
        else:
            ev.refund_amount = total_amount
    if parsed["refund_reason"] and not (only_fill_empty and (ev.refund_reason or "").strip()):
        ev.refund_reason = parsed["refund_reason"]
    if parsed["return_logistics_no"]:
        ev.return_logistics_no = parsed["return_logistics_no"]
        ev.return_carrier = parsed["return_carrier"]
    goods_count = sum(
        (client.parse_refund(x)["return_goods_count"] for x in rows), Decimal("0")
    )
    if goods_count > 0 or not only_fill_empty:
        ev.return_goods_count = goods_count
    ev.wdt_raw = json.dumps(rows, ensure_ascii=False, default=str)
    ev.wdt_synced_at = datetime.now()

    # 商品级退回明细：所有退款单的明细合并（按 spec_no 去重累加数量）
    ev.refund_items.clear()
    merged: dict[str, dict] = {}
    for row in rows:
        for item in client.parse_refund_items(row):
            key = item["spec_no"] or item["goods_name"]
            if key in merged:
                for f in ("order_num", "refund_num", "stockin_num"):
                    merged[key][f] += item[f]
            else:
                merged[key] = item
    ev.refund_items.extend(RefundItem(**x) for x in merged.values())
    return total_amount


@router.post("/events/{code}/sync-wdt")
def sync_wdt(code: str, ident: Identity = Depends(require_editor), db: Session = Depends(get_db)):
    """按订单号/快递单号自动拉退款单+退回明细，落快照。"""
    actor = ident.email
    ev = _get(db, code)
    client = WangdianClient()
    if not client.configured:
        raise HTTPException(400, "旺店通凭证未配置")
    # 发货信息独立于退款信息查询。一个订单已经退款时，仍然必须查销售出库单，
    # 否则“发货快递公司/发货单号”不会被带回理赔台账。
    stockouts: list = []
    so_err = ""
    try:
        stockouts = client.find_sales_stockouts(
            src_tid=ev.wdt_tid or "", logistics_no=ev.ship_logistics_no or ""
        )
    except (WdtError, ValueError) as exc:
        so_err = str(exc)
        logging.getLogger("claims").warning("出库单查询失败 tid=%s: %s", ev.wdt_tid, so_err)

    shipment = None
    filled: list[str] = []
    if stockouts:
        shipment = client.parse_sales_stockout(stockouts[0])
        _apply_shipment(ev, shipment)
        if not (ev.wdt_tid or "").strip() and shipment.get("source_tid"):
            ev.wdt_tid = shipment["source_tid"]
            filled.append(f"原始订单号 {shipment['source_tid']}")
        raw_shipment = shipment.get("raw") or {}
        if not (ev.shop_name or "").strip() and raw_shipment.get("shop_name"):
            ev.shop_name = str(raw_shipment["shop_name"])
            ev.platform = platform_of(ev.shop_name) or ev.platform
        if not (ev.ship_logistics_no or "").strip() and shipment["logistics_no"]:
            ev.ship_logistics_no = shipment["logistics_no"]
            filled.append(f"发货物流 {shipment['logistics_no']}")
        if not (ev.ship_carrier or "").strip() and shipment["carrier"]:
            ev.ship_carrier = shipment["carrier"]
            filled.append(f"快递公司 {shipment['carrier']}")
    # 发货运单与退货运单不是同一字段；先解析原始订单，再查询售后。
    rows: list = []
    if ev.wdt_tid or ev.return_logistics_no:
        try:
            rows = client.find_refunds(tid=ev.wdt_tid or "",
                                      logistics_no="" if ev.wdt_tid else ev.return_logistics_no)
        except (WdtError, ValueError) as exc:
            raise HTTPException(502, str(exc)) from exc
    if not rows:
        # 退款单查不到 ≠ 异常：订单可能尚未发货、退款尚未发生（用户确认的逻辑）。
        # 降级查销售出库单：只要货发过，就能带出发货物流/时间，理赔举证可用。
        if stockouts:
            ev.wdt_synced_at = datetime.now()
            db.commit()
            db.refresh(ev)
            audit_action(db, ev, actor, "sync", field="-",
                         new_value=f"暂无退款单；查到发货记录 {shipment['stockout_no']}"
                                   f"（{shipment['consign_time'] or '时间未知'}）")
            db.commit()
            return {
                "ok": True,
                "stage": "stockout",
                "event": serialize_event(ev),
                "message": (
                    f"旺店通暂无退款单（退款尚未发生属正常）。"
                    f"查到发货记录：出库单 {shipment['stockout_no']}"
                    + (f"，{shipment['carrier']}" if shipment["carrier"] else "")
                    + (f"，物流 {shipment['logistics_no']}" if shipment["logistics_no"] else "")
                    + (f"，发货 {shipment['consign_time']}" if shipment["consign_time"] else "")
                    + ("。已自动回填发货物流单号" + "、".join(filled) if filled else "")
                ),
            }
        return {
            "ok": False,
            "stage": "none",
            "message": (
                "旺店通接口未返回匹配的订单/发货信息。"
                + ("当前普通接口不返回淘宝、天猫订单；这些订单需配置奇门自定义接口。"
                   if not client.qimen_configured else "请核对旺店通店铺授权和查询号码。")
            )
            + (f"［出库单查询异常：{so_err}］" if so_err else ""),
        }

    _persist_refunds(db, ev, client, rows)
    if shipment:
        # 即使退款接口命中，也把销售出库的快递事实写入事件快照。
        # 已人工填写的值优先，避免同步覆盖人工修正。
        if not (ev.ship_logistics_no or "").strip() and shipment["logistics_no"]:
            ev.ship_logistics_no = shipment["logistics_no"]
        if not (ev.ship_carrier or "").strip() and shipment["carrier"]:
            ev.ship_carrier = shipment["carrier"]
        if not (ev.wdt_tid or "").strip() and shipment.get("source_tid"):
            ev.wdt_tid = shipment["source_tid"]

    # 编号依据 = 订单号（用户确认）。仍是 CLM 临时编号且订单号可用 → 改为订单号
    renamed = False
    tid = (ev.wdt_tid or "").strip()
    if tid and ev.code != tid and ev.code.startswith("CLM-"):
        dup = db.scalar(select(ClaimEvent).where(ClaimEvent.code == tid))
        if dup is None:
            old_code = ev.code
            ev.code = tid
            renamed = True
            ev.note = ("[编号] 由 " + old_code + " 改用订单号 " + tid).strip()

    db.commit()
    db.refresh(ev)
    audit_action(db, ev, actor, "sync", field="-",
                 new_value=f"退款{len(rows)}笔，顾客已赔合计 {ev.refund_amount}")
    db.commit()
    result = {"ok": True, "stage": "refund", "event": serialize_event(ev), "refund_count": len(rows)}
    if shipment:
        result["shipment"] = shipment
        result["message"] = (
            f"已同步退款 {len(rows)} 笔；"
            + (f"发货快递 {shipment['carrier']}，" if shipment["carrier"] else "")
            + (f"运单号 {shipment['logistics_no']}" if shipment["logistics_no"] else "")
        ).rstrip("，")
    elif so_err:
        result["shipment_message"] = f"发货信息查询失败：{so_err}"
    if renamed:
        result["renamed_from"] = old_code
        result["message"] = (result.get("message", "") + "；" if result.get("message") else "") + "编号已改用订单号 " + tid
    elif len(rows) > 1:
        result["message"] = (result.get("message", "") + "；" if result.get("message") else "") + "该订单共 " + str(len(rows)) + " 笔退款，已全部合入本登记"
    return result


@router.post("/events/{code}/check-stockin")
def check_stockin(code: str, ident: Identity = Depends(require_editor), db: Session = Depends(get_db)):
    """一键查退货入库：状态/仓库（正品vs次品）/商品档案/效期。"""
    actor = ident.email
    ev = _get(db, code)
    if not (ev.wdt_tid or "").strip() and not (ev.return_logistics_no or "").strip():
        raise HTTPException(400, "缺少原始订单号或退货单号，不能判断入库情况，请先完成旺店通核验")
    client = WangdianClient()
    if not client.configured:
        raise HTTPException(400, "旺店通凭证未配置")
    try:
        rows = client.search_refund_stockins(
            order_no=ev.wdt_tid, logistics_no=ev.return_logistics_no
        )
    except (WdtError, ValueError) as exc:
        raise HTTPException(502, str(exc)) from exc

    if not rows:
        ev.stockin_checked_at = datetime.now()
        db.commit()
        db.refresh(ev)
        audit_action(db, ev, actor, "stockin", field="-", new_value="未查到入库")
        db.commit()
        return {"ok": True, "found": ev.stockin_found, "message": "当前查询窗口未匹配到入库记录，保留历史查询证据，不代表从未入库", "event": serialize_event(ev)}

    row = max(rows, key=lambda item: str(item.get("stockin_time") or ""))
    parsed = client.parse_stockin(row)
    ev.stockin_found = True
    ev.stockin_status = parsed["stockin_status"]
    ev.stockin_warehouse = "、".join(dict.fromkeys(client.parse_stockin(r)["stockin_warehouse"] for r in rows))
    ev.stockin_is_defective = any(client.parse_stockin(r)["stockin_is_defective"] for r in rows)
    ev.stockin_time = parsed["stockin_time"]
    ev.stockin_raw = json.dumps(rows, ensure_ascii=False, default=str)
    ev.stockin_checked_at = datetime.now()

    ev.stockin_items.clear()
    for matched_row in rows:
        for item in client.parse_stockin_items(matched_row):
            ev.stockin_items.append(StockinItem(**item))
    db.commit()
    db.refresh(ev)
    audit_action(db, ev, actor, "stockin", field="-",
                 new_value=f"{'次品仓' if ev.stockin_is_defective else '已入库'} status={ev.stockin_status}")
    db.commit()
    return {"ok": True, "found": True, "event": serialize_event(ev)}


@router.get("/wdt/backfill-returns/candidates")
def backfill_return_candidates(
    month: str = Query(default=""),
    start: str = Query(default=""),
    end: str = Query(default=""),
    db: Session = Depends(get_db),
):
    """列出「退货物流 / 退货单号」仍为空的记录（只读，不调旺店通）。

    这两个字段只来自旺店通退款单（return_logistics_name / return_logistics_no），
    而详情页原先只查发货信息、不查退款单，所以历史数据全为空。
    """
    q = _apply_filters(select(ClaimEvent).order_by(ClaimEvent.event_date), month, start, end)
    q = q.where(
        (func.coalesce(ClaimEvent.return_logistics_no, "") == "")
        | (func.coalesce(ClaimEvent.return_carrier, "") == "")
    )
    evs = list(db.scalars(q).unique().all())
    return {
        "total": len(evs),
        "codes": [ev.code for ev in evs],
        "items": [{"code": ev.code, "wdt_tid": ev.wdt_tid,
                   "ship_logistics_no": ev.ship_logistics_no} for ev in evs[:200]],
    }


class BackfillIn(BaseModel):
    codes: list[str]
    dry_run: bool = False


@router.post("/wdt/backfill-returns")
def backfill_returns(
    payload: BackfillIn,
    ident: Identity = Depends(require_admin_identity),
    db: Session = Depends(get_db),
):
    """按给定编号批量补取退货物流/退货单号/退货件数。

    只填空值（only_fill_empty），不覆盖人工核对过的金额与原因。
    前端分批调用，避免单次请求超时。
    """
    actor = ident.email
    client = WangdianClient()
    if not client.configured:
        raise HTTPException(400, "旺店通凭证未配置")

    filled: list[dict] = []
    no_refund: list[str] = []
    errors: list[dict] = []

    for code in payload.codes:
        ev = db.scalar(select(ClaimEvent).where(ClaimEvent.code == code))
        if ev is None:
            errors.append({"code": code, "error": "记录不存在"})
            continue
        tid = (ev.wdt_tid or "").strip()
        lno = (ev.ship_logistics_no or "").strip()
        if not tid and not lno:
            errors.append({"code": code, "error": "缺少订单号与快递单号"})
            continue
        try:
            rows = client.find_refunds(tid=tid, logistics_no="" if tid else lno)
        except (WdtError, ValueError) as exc:
            errors.append({"code": code, "error": str(exc)})
            continue
        if not rows:
            no_refund.append(code)
            continue

        parsed = client.parse_refund(rows[0])
        if not (parsed["return_logistics_no"] or parsed["return_carrier"]):
            no_refund.append(code)
            continue

        filled.append({
            "code": code,
            "wdt_tid": ev.wdt_tid,
            "return_carrier": parsed["return_carrier"],
            "return_logistics_no": parsed["return_logistics_no"],
            "return_goods_count": str(parsed["return_goods_count"]),
            "refund_no": parsed["wdt_refund_no"],
        })
        if not payload.dry_run:
            _persist_refunds(db, ev, client, rows, only_fill_empty=True)
            audit_action(
                db, ev, actor, "sync", field="return_logistics_no",
                new_value=f"批量补取退货信息：{parsed['return_carrier']} {parsed['return_logistics_no']}",
            )

    if not payload.dry_run:
        db.commit()

    return {
        "ok": True,
        "dry_run": payload.dry_run,
        "scanned": len(payload.codes),
        "filled": len(filled),
        "no_refund": len(no_refund),
        "errors": len(errors),
        "samples": filled[:20],
        "no_refund_codes": no_refund[:20],
        "error_samples": errors[:10],
    }


@router.get("/wdt/test")
def wdt_test():
    client = WangdianClient()
    if not client.configured:
        raise HTTPException(400, "旺店通凭证未配置")
    try:
        info = client.test_connection()
        return {"ok": True, "message": "连接成功", "shop_count": info["shop_count"]}
    except (WdtError, Exception) as exc:  # noqa: BLE001
        raise HTTPException(502, str(exc)) from exc


# ---------------------------------------------------------------- 附件


ALLOWED_SUFFIX = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".pdf", ".txt", ".csv", ".xlsx"}
EVIDENCE_PARTIES = {"warehouse", "merchant", "express", "buyer", "unclassified"}


class EvidencePartyIn(BaseModel):
    source_party: str


@router.put("/evidence/{eid}/source")
def classify_evidence(eid: int, payload: EvidencePartyIn, ident: Identity = Depends(require_editor),
                      db: Session = Depends(get_db)):
    if payload.source_party not in EVIDENCE_PARTIES:
        raise HTTPException(400, "请选择有效的凭证来源")
    evidence = db.get(Evidence, eid)
    if evidence is None:
        raise HTTPException(404, "附件不存在")
    parent = db.get(ClaimEvent, evidence.event_id)
    if parent is not None:
        _get(db, parent.code)
    evidence.source_party = payload.source_party
    db.commit()
    return {"ok": True}


@router.post("/events/{code}/evidence")
async def upload_evidence(
    code: str, file: UploadFile = File(...),
    kind: str = Form(default="other"), caption: str = Form(default=""),
    source_party: str = Form(default="unclassified"),
    ident: Identity = Depends(require_editor),
    db: Session = Depends(get_db),
):
    ev = _get(db, code)
    if source_party not in EVIDENCE_PARTIES:
        raise HTTPException(400, "请选择有效的凭证来源")
    suffix = Path(file.filename or "").suffix.lower()
    if suffix and suffix not in ALLOWED_SUFFIX:
        raise HTTPException(400, f"不支持的文件类型：{suffix}")
    root = Path(settings.upload_dir)
    root.mkdir(parents=True, exist_ok=True)
    stored = f"{ev.code}_{uuid.uuid4().hex[:8]}{suffix}"
    dest = root / stored
    size = 0
    limit = settings.max_upload_mb * 1024 * 1024
    with dest.open("wb") as fh:
        while chunk := await file.read(1024 * 1024):
            size += len(chunk)
            if size > limit:
                fh.close()
                dest.unlink(missing_ok=True)
                raise HTTPException(413, f"文件超过 {settings.max_upload_mb}MB 上限")
            fh.write(chunk)
    db.add(
        Evidence(event_id=ev.id, kind=kind, filename=file.filename or stored,
                 stored_path=stored, content_type=file.content_type or "",
                 size_bytes=size, caption=caption, source_party=source_party)
    )
    db.commit()
    db.refresh(ev)
    return {"ok": True, "evidence_count": len(ev.evidence), "event": serialize_event(ev)}


@router.get("/evidence/{eid}/download")
def download_evidence(eid: int, db: Session = Depends(get_db)):
    ev = db.get(Evidence, eid)
    if ev is None:
        raise HTTPException(404, "附件不存在")
    p = Path(settings.upload_dir) / ev.stored_path
    if not p.exists():
        raise HTTPException(410, f"附件文件已丢失：{ev.stored_path}")
    return FileResponse(p, filename=ev.filename or p.name)


@router.delete("/evidence/{eid}")
def delete_evidence(eid: int, ident: Identity = Depends(require_editor), db: Session = Depends(get_db)):
    ev = db.get(Evidence, eid)
    if ev is None:
        raise HTTPException(404, "附件不存在")
    parent = db.get(ClaimEvent, ev.event_id)
    if parent is not None:
        _get(db, parent.code)
    p = Path(settings.upload_dir) / ev.stored_path
    if p.exists():
        p.unlink()
    db.delete(ev)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- 导入导出


@router.post("/import/legacy")
async def import_legacy_file(file: UploadFile = File(...),
                             ident: Identity = Depends(require_admin_identity),
                             db: Session = Depends(get_db)):
    content = await file.read()
    if not content:
        raise HTTPException(400, "文件为空")
    from .services.legacy_import import import_legacy  # noqa: PLC0415

    try:
        return import_legacy(db, content, verify_shipment=_lookup_shipment,
                             apply_shipment=_apply_shipment, validate_amounts=_validate_amounts)
    except Exception as exc:  # noqa: BLE001
        db.rollback()
        raise HTTPException(400, f"导入失败：{exc}") from exc


@router.get("/export/csv")
def export_csv_file(
    db: Session = Depends(get_db), month: str = Query(default=""),
    start: str = Query(default=""), end: str = Query(default=""),
):
    """导出 CSV（供程序对接）。

    Excel 直接打开 CSV 时，会把超过 15 位的纯数字单号（如原始订单号）转成科学计数法
    并丢掉末几位——CSV 不带类型信息，这里挡不住；给人看的 Excel 请用 /export/xlsx。
    枚举（售后原因 / 售后类型 / 状态）输出中文，不再是 loss / damage / original_return。
    """
    q = _apply_filters(select(ClaimEvent).order_by(ClaimEvent.event_date), month, start, end)
    evs = list(db.scalars(q).unique().all())
    lines = [",".join(EXPORT_HEADERS)]
    for row in export_rows(evs):
        lines.append(",".join(f'"{str(c).replace(chr(34), chr(34)*2)}"' for c in row))
    return Response(
        content="﻿" + "\r\n".join(lines),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=\"claims.csv\"; filename*=UTF-8''{quote('售后登记表.csv')}"},
    )


@router.get("/export/xlsx")
def export_xlsx_file(
    db: Session = Depends(get_db), month: str = Query(default=""),
    start: str = Query(default=""), end: str = Query(default=""),
):
    """导出单表 Excel 明细：字段与 CSV 完全一致，但单号列按文本写入。

    这样 Excel 打开不会把长订单号变成 1.10211E+17，也不会丢末几位；
    枚举（售后原因 / 售后类型 / 状态）输出中文。
    """
    q = _apply_filters(select(ClaimEvent).order_by(ClaimEvent.event_date), month, start, end)
    evs = list(db.scalars(q).unique().all())
    fname = quote("售后登记表_明细.xlsx")
    return Response(
        content=build_export_workbook(evs),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=\"claims-export.xlsx\"; filename*=UTF-8''{fname}"
        },
    )


@router.get("/export/template")
def import_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "导入模板"
    cols = [
        "事件编号", "日期", "店铺名称", "订单号", "快递公司", "快递单号", "售后类型", "售后内容",
        "仓库赔付金额", "责任方", "实际到手", "状态", "仓库回复", "货主备注", "备注",
    ]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E8E4")
    ws.append([
        "CLM20260801-001", "2026-08-01", "梦奇思-抖音-梦奇思零食专营店", "6954640341931005538",
        "京东快递-青浦仓", "JDAZ26408198637", "丢件", "丢件-仓库没有做入库",
        "24.9", "仓库", "", "已提交", "麻烦提供售后证明", "已提供", "客户已确认丢件",
    ])
    for i in range(1, len(cols) + 1):
        from openpyxl.utils import get_column_letter  # noqa: PLC0415

        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="claims-import-template.xlsx"'},
    )


# ---------------------------------------------------------------- 报表（电子，非打印）

@router.get("/report/xlsx")
def report_xlsx(
    db: Session = Depends(get_db), month: str = Query(default=""),
    start: str = Query(default=""), end: str = Query(default=""),
):
    """导出多 sheet 格式化 Excel 报表：汇总 / 明细 / 按责任方 / 逾期在途。"""
    q = _apply_filters(select(ClaimEvent).order_by(ClaimEvent.event_date), month, start, end)
    evs = list(db.scalars(q).unique().all())
    data = summarize(evs)
    xlsx = build_report_workbook(evs, data)
    # HTTP 头只能承载 latin-1，中文文件名需 RFC 5987 编码，否则 starlette 抛 500。
    fname = quote("售后登记表.xlsx")
    return Response(
        content=xlsx,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"claims-report.xlsx\"; filename*=UTF-8''{fname}"},
    )


@router.post("/report/import")
async def report_import(
    file: UploadFile = File(...), ident: Identity = Depends(require_editor),
    db: Session = Depends(get_db),
):
    """回填导入：读取报表「明细」sheet，按编号更新 实收金额 / 状态 / 仓库工单号。

    导出 → 在 Excel 里填好回款 → 再导入回来批量更新，形成电子闭环。
    仅更新被改动过的行；未改的行原样回写不产生审计噪声。
    """
    actor = ident.email
    # 回填字段属于仓库口径，渠道不符时直接拒绝，避免客服导错表把回款覆盖。
    check_fields(ident, ["received_amount", "status", "warehouse_ref"])
    content = await file.read()
    if not content:
        raise HTTPException(400, "文件为空")
    try:
        updates = parse_report_import(content)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"解析报表失败：{exc}")
    updated = 0
    notfound = 0
    for u in updates:
        identity_filter = ClaimEvent.wdt_tid == u["wdt_tid"] if "wdt_tid" in u else ClaimEvent.code == u["code"]
        matches = list(db.scalars(select(ClaimEvent).where(identity_filter,
            ~ClaimEvent.id.in_(merged_event_ids())).limit(2)).all())
        if len(matches) > 1:
            raise HTTPException(409, "原始订单号存在多条登记，请先核对合并后再回填")
        ev = matches[0] if matches else None
        if ev is None:
            notfound += 1
            continue
        before = {a: getattr(ev, a) for a, _ in AUDIT_FIELDS}
        _fill(ev, EventIn(**{key: u[key] for key in ("received_amount", "status", "warehouse_ref") if key in u}))
        audit_diff(db, ev, actor, before)
        updated += 1
    db.commit()
    return {"ok": True, "updated": updated, "notfound": notfound, "total": len(updates)}


# ---------------------------------------------------------------- 企业微信工单建单
#
# 走 open_router（不经用户登录体系），用独立 X-Bot-Token 保护；disable_auth 开发期放行。
# 设计：建单权只交给 @ 触发（wecom_msg_id 幂等）；同订单+同分类+24h 内重复 → 合并归集。

class WorkOrderIn(BaseModel):
    customer_name: str = ""
    customer_wecom_id: str = ""
    channel: str = ""
    related_order_no: str = ""
    category: str = "other"
    priority: str = "medium"
    summary: str = ""
    source_message: str = ""   # 原始 @ 消息内容，写首条 timeline
    wecom_msg_id: str = ""      # 企业微信原始消息 id（幂等去重）
    source: str = "wecom_group"


class WorkOrderUpdate(BaseModel):
    status: str = ""        # 目标状态（英文枚举值或中文）
    priority: str = ""      # 优先级（英文枚举值或中文）
    owner: str = ""         # 负责人
    summary: str = ""       # 摘要覆盖
    note: str = ""          # 人工备注，追加一条 MANUAL timeline
    inferred_status: str = ""  # 预留：AI 从群聊推断的状态（扫描层用）


def _auth_user_or_bot(request: Request) -> str:
    """工单接口鉴权：用户登录 token（人用页面）或 bot token（bot 服务端调用）二选一。

    返回操作人标识：dev(开发免登录) / 用户邮箱 / "bot"。
    生产环境（disable_auth=False）下两者皆无则 401，绝不裸奔。
    """
    if settings.disable_auth:
        return "dev"
    # 1) 用户登录 token：Authorization: Bearer <token>
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        email = verify_token(auth[7:].strip())
        if email:
            return email
    # 2) bot token：X-Bot-Token（服务端调用，与 bot 侧 WECOM_BOT_TOKEN 一致）
    expected = settings.wecom_bot_token
    if expected and request.headers.get("X-Bot-Token") == expected:
        return "bot"
    raise HTTPException(401, "未登录或 token 无效")


# 工单分类/优先级：bot 可能传中文（破损/高），统一收口到英文枚举值。
WOCAT_ALIAS = {"破损": "damage", "damage": "damage", "漏发": "missing", "missing": "missing",
               "退货": "return", "return": "return", "换货": "exchange", "exchange": "exchange",
               "错发": "wrong", "wrong": "wrong", "其他": "other", "other": "other"}
WOPRI_ALIAS = {"高": "high", "high": "high", "中": "medium", "medium": "medium",
               "低": "low", "low": "low"}


def _enum_wo_cat(v: str) -> WorkOrderCategory:
    key = (v or "other").strip()
    return WorkOrderCategory(WOCAT_ALIAS.get(key, WOCAT_ALIAS.get(key.lower(), "other")))


def _enum_wo_pri(v: str) -> WorkOrderPriority:
    key = (v or "中").strip()
    return WorkOrderPriority(WOPRI_ALIAS.get(key, WOPRI_ALIAS.get(key.lower(), "medium")))


def _enum_wo_src(v: str) -> WorkOrderSource:
    try:
        return WorkOrderSource((v or "wecom_group").strip().lower())
    except ValueError:
        return WorkOrderSource.WECOM_GROUP


WOSTAT_ALIAS = {"待处理": "pending", "pending": "pending", "处理中": "processing",
                "processing": "processing", "待客户": "waiting", "waiting": "waiting",
                "已解决": "resolved", "resolved": "resolved", "已关闭": "closed",
                "closed": "closed", "超时": "overdue", "overdue": "overdue"}


def _enum_wo_status(v: str) -> WorkOrderStatus:
    key = (v or "pending").strip()
    return WorkOrderStatus(WOSTAT_ALIAS.get(key, WOSTAT_ALIAS.get(key.lower(), "pending")))


def _next_wo_code(db: Session) -> str:
    today = date.today().strftime("%Y%m%d")
    prefix = f"WO-{today}-"
    last = db.scalar(
        select(WorkOrder.code).where(WorkOrder.code.like(prefix + "%"))
        .order_by(WorkOrder.id.desc()).limit(1)
    )
    n = 1
    if last:
        try:
            n = int(last[len(prefix):]) + 1
        except ValueError:
            n = 1
    return f"{prefix}{n:03d}"


def _serialize_wo(wo: WorkOrder) -> dict:
    return {
        "id": wo.id, "code": wo.code, "source": wo.source.value,
        "customer_name": wo.customer_name, "customer_wecom_id": wo.customer_wecom_id,
        "channel": wo.channel, "related_order_no": wo.related_order_no,
        "category": wo.category.value, "priority": wo.priority.value,
        "status": wo.status.value, "owner": wo.owner, "summary": wo.summary,
        "wecom_msg_id": wo.wecom_msg_id,
        "created_at": wo.created_at.isoformat() if wo.created_at else "",
        "timeline_count": len(wo.timelines),
    }


@open_router.post("/work_orders")
def create_work_order(payload: WorkOrderIn, request: Request, db: Session = Depends(get_db)):
    """企业微信 bot 调用：@ 触发登记工单。

    三档去重/归集：
      1) 同一 wecom_msg_id 已建过 → 直接返回（幂等，避免 @ 机器人被重复触发刷单）；
      2) 同订单号 + 同分类 + 24h 内活跃工单 → 把新消息合并进现有单（补 timeline，不新建）；
      3) 否则新建工单 + 首条 timeline（机器人建单）。
    """
    actor = _auth_user_or_bot(request)
    # 1) 幂等
    if payload.wecom_msg_id:
        dup = db.scalar(select(WorkOrder).where(WorkOrder.wecom_msg_id == payload.wecom_msg_id).limit(1))
        if dup:
            return {"ok": True, "dup": True, "order": _serialize_wo(dup), "message": "该消息已建单"}
    # 2) 防重合并
    if payload.related_order_no:
        merged = db.scalar(
            select(WorkOrder).where(
                WorkOrder.related_order_no == payload.related_order_no,
                WorkOrder.category == _enum_wo_cat(payload.category),
                WorkOrder.status.in_([
                    WorkOrderStatus.PENDING, WorkOrderStatus.PROCESSING, WorkOrderStatus.WAITING_CUSTOMER
                ]),
                WorkOrder.created_at >= datetime.now() - timedelta(hours=24),
            ).order_by(WorkOrder.id.desc()).limit(1)
        )
        if merged:
            merged.timelines.append(WorkOrderTimeline(
                source=TimelineSource.WECOM_SCAN, author=payload.customer_name,
                content=payload.source_message or payload.summary,
            ))
            db.commit()
            return {"ok": True, "merged": True, "order": _serialize_wo(merged), "message": "已合并到现有工单"}
    # 3) 新建
    wo = WorkOrder(
        code=_next_wo_code(db), source=_enum_wo_src(payload.source),
        customer_name=payload.customer_name, customer_wecom_id=payload.customer_wecom_id,
        channel=payload.channel, related_order_no=payload.related_order_no,
        category=_enum_wo_cat(payload.category), priority=_enum_wo_pri(payload.priority),
        status=WorkOrderStatus.PENDING, summary=payload.summary, wecom_msg_id=payload.wecom_msg_id,
    )
    author = "售后助手" if (actor == "bot" or payload.source == "wecom_group") else (actor or "运营")
    wo.timelines.append(WorkOrderTimeline(
        source=TimelineSource.BOT_CREATED, author=author,
        content=payload.source_message or payload.summary,
    ))
    db.add(wo)
    db.commit()
    db.refresh(wo)
    return {"ok": True, "order": _serialize_wo(wo), "message": "已新建工单"}


def _serialize_tl(tl: WorkOrderTimeline) -> dict:
    return {
        "id": tl.id, "source": tl.source.value,
        "source_label": TL_SRC_LABEL.get(tl.source.value, tl.source.value),
        "author": tl.author, "content": tl.content,
        "inferred_status": tl.inferred_status,
        "created_at": tl.created_at.isoformat() if tl.created_at else "",
    }


def _serialize_wo_detail(wo: WorkOrder) -> dict:
    d = _serialize_wo(wo)
    d["category_label"] = WOCAT_LABEL.get(wo.category.value, wo.category.value)
    d["priority_label"] = WOPRI_LABEL.get(wo.priority.value, wo.priority.value)
    d["status_label"] = WOSTAT_LABEL.get(wo.status.value, wo.status.value)
    d["timelines"] = [_serialize_tl(t) for t in wo.timelines]
    return d


@open_router.get("/work_orders")
def list_work_orders(
    request: Request,
    status: str = "", category: str = "", priority: str = "", owner: str = "",
    keyword: str = "", source: str = "", limit: int = 50, offset: int = 0,
    db: Session = Depends(get_db),
):
    """工单看板列表（运营/客服与仓库端共用，前端按角色做字段/范围过滤）。"""
    _auth_user_or_bot(request)
    q = select(WorkOrder)
    if status:
        vals = [_enum_wo_status(s) for s in status.split(",") if s.strip()]
        if vals:
            q = q.where(WorkOrder.status.in_(vals))
    if category:
        cats = [_enum_wo_cat(c).value for c in category.split(",") if c.strip()]
        if cats:
            q = q.where(WorkOrder.category.in_(cats))
    if priority:
        pris = [_enum_wo_pri(p).value for p in priority.split(",") if p.strip()]
        if pris:
            q = q.where(WorkOrder.priority.in_(pris))
    if owner:
        q = q.where(WorkOrder.owner.ilike(f"%{owner}%"))
    if source:
        q = q.where(WorkOrder.source == _enum_wo_src(source))
    if keyword:
        kw = f"%{keyword.strip()}%"
        q = q.where(or_(
            WorkOrder.related_order_no.ilike(kw), WorkOrder.customer_name.ilike(kw),
            WorkOrder.code.ilike(kw), WorkOrder.summary.ilike(kw),
        ))
    total = db.scalar(select(func.count()).select_from(q.subquery())) or 0
    rows = db.scalars(
        q.order_by(WorkOrder.created_at.desc()).limit(limit).offset(offset)
    ).all()
    return {"ok": True, "total": total, "rows": [_serialize_wo(r) for r in rows]}


@open_router.get("/work_orders/stats")
def work_order_stats(request: Request, db: Session = Depends(get_db)):
    """看板指标卡：待处理 / 处理中 / 今日已解决 / 超时预警。"""
    _auth_user_or_bot(request)
    today = date.today()
    pending = db.scalar(
        select(func.count()).select_from(WorkOrder).where(WorkOrder.status == WorkOrderStatus.PENDING)
    ) or 0
    processing = db.scalar(
        select(func.count()).select_from(WorkOrder).where(
            WorkOrder.status.in_([WorkOrderStatus.PROCESSING, WorkOrderStatus.WAITING_CUSTOMER])
        )
    ) or 0
    resolved_today = db.scalar(
        select(func.count()).select_from(WorkOrder).where(
            WorkOrder.status.in_([WorkOrderStatus.RESOLVED, WorkOrderStatus.CLOSED]),
            WorkOrder.updated_at >= datetime(today.year, today.month, today.day),
        )
    ) or 0
    overdue = db.scalar(
        select(func.count()).select_from(WorkOrder).where(WorkOrder.status == WorkOrderStatus.OVERDUE)
    ) or 0
    return {"ok": True, "stats": {
        "pending": pending, "processing": processing,
        "resolved_today": resolved_today, "overdue": overdue,
    }}


@open_router.get("/work_orders/{wo_id}")
def get_work_order(wo_id: int, request: Request, db: Session = Depends(get_db)):
    _auth_user_or_bot(request)
    wo = db.get(WorkOrder, wo_id)
    if not wo:
        raise HTTPException(404, "工单不存在")
    return {"ok": True, "order": _serialize_wo_detail(wo)}


@open_router.patch("/work_orders/{wo_id}")
def update_work_order(wo_id: int, payload: WorkOrderUpdate, request: Request, db: Session = Depends(get_db)):
    """人工覆盖：改状态/优先级/负责人/摘要，或追加一条内部备注（MANUAL timeline）。

    状态/优先级的取值兼容中文（待处理/高）。note 非空时自动挂时间线，保证可追溯。
    """
    actor = _auth_user_or_bot(request)
    wo = db.get(WorkOrder, wo_id)
    if not wo:
        raise HTTPException(404, "工单不存在")
    if payload.status:
        wo.status = _enum_wo_status(payload.status)
    if payload.priority:
        wo.priority = _enum_wo_pri(payload.priority)
    if payload.owner:
        wo.owner = payload.owner
    if payload.summary:
        wo.summary = payload.summary
    if payload.note:
        wo.timelines.append(WorkOrderTimeline(
            source=TimelineSource.MANUAL, author=payload.owner or actor or "运营", content=payload.note,
        ))
    db.commit()
    db.refresh(wo)
    return {"ok": True, "order": _serialize_wo_detail(wo), "message": "已更新"}
